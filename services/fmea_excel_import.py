"""Read the company's 갑지 analysis table without importing document headers or writing data."""
import io
import re
from datetime import date, datetime
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_SHEET_ROWS = 2000
MAX_ANALYSES = 500


class FmeaImportError(ValueError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("엑셀 불러오기를 취소했습니다. 기존 분석행은 변경되지 않았습니다.")


def compact(value):
    return re.sub(r"\s+", "", str(value or ""))


def parse_company_fmea(content, steps):
    errors = []

    def fail(cell, message):
        error = {"cell": f"갑지!{cell}" if cell else "파일", "message": message}
        if error not in errors:
            errors.append(error)

    if not content or len(content) > MAX_FILE_BYTES:
        raise FmeaImportError([{"cell": "파일", "message": "10MB 이하의 비어 있지 않은 .xlsx 파일을 선택해 주세요."}])
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(x.file_size for x in entries) > MAX_EXPANDED_BYTES:
                raise ValueError("expanded size")
            if any(x.filename.endswith('vbaProject.bin') for x in entries):
                raise ValueError("macros")
        workbook = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError, OSError, EOFError, SyntaxError, TypeError, IndexError, NotImplementedError, RuntimeError):
        raise FmeaImportError([{"cell": "파일", "message": "파일이 손상되었거나 지원하지 않는 엑셀 형식입니다. 매크로 없는 .xlsx 파일을 사용해 주세요."}]) from None
    try:
        if "갑지" not in workbook.sheetnames:
            raise FmeaImportError([{"cell": "파일", "message": "'갑지' 시트가 없습니다. 회사 공정 FMEA 양식을 사용해 주세요."}])
        sheet = workbook["갑지"]
        if sheet.max_row > MAX_SHEET_ROWS or sheet.max_column > 40:
            raise FmeaImportError([{"cell": "파일", "message": "갑지의 행·열 범위가 지원하는 회사 양식을 초과합니다."}])
        headers = [r for r in range(1, min(sheet.max_row, 100) + 1)
                   if compact(sheet.cell(r, 1).value) == "공정의기능"
                   and compact(sheet.cell(r, 2).value) == "잠재적고장형태"]
        if len(headers) != 1:
            raise FmeaImportError([{"cell": "파일", "message": "갑지에서 분석표 항목명을 하나로 확인할 수 없습니다."}])
        header = headers[0]
        labels = {1: "공정의기능", 2: "잠재적고장형태", 3: "고장의잠재적영향", 4: "심각도",
                  5: "특별특성", 6: "고장의잠재적원인", 8: "발생도", 9: "현공정관리",
                  11: "검출도", 12: "R.P.N.", 13: "권고조치사항", 15: "완료예정일", 16: "조치결과"}
        for col, expected in labels.items():
            if compact(sheet.cell(header, col).value) != expected:
                fail(f"{get_column_letter(col)}{header}", f"항목명이 '{expected}'와 일치하지 않습니다.")
        for col, expected in {16: "조치내용", 17: "심각도", 18: "발생도", 19: "검출도", 20: "R.P.N."}.items():
            if compact(sheet.cell(header + 1, col).value) != expected:
                fail(f"{get_column_letter(col)}{header + 1}", f"조치결과 항목 '{expected}'를 확인해 주세요.")
        for col, expected in {9: "예방", 10: "검출"}.items():
            if compact(sheet.cell(header + 2, col).value) != expected:
                fail(f"{get_column_letter(col)}{header + 2}", f"현공정 관리의 '{expected}' 열을 확인해 주세요.")
        if errors:
            raise FmeaImportError(errors)
        start = header + 3
        footers = [r for r in range(start, sheet.max_row + 1)
                   if compact(sheet.cell(r, 1).value).startswith("RPN(")]
        end = footers[0] - 1 if footers else sheet.max_row
        # Resolve only actual merged cells; never fill unrelated blank cells downwards.
        anchors = {}
        for merged in sheet.merged_cells.ranges:
            if merged.max_row < start or merged.min_row > end:
                continue
            if merged.min_row < start or merged.max_row > end:
                fail(str(merged), "병합 셀이 분석표 경계를 넘어갑니다.")
                continue
            if merged.min_col != merged.max_col and (merged.min_col, merged.max_col) not in {(6, 7), (13, 14)}:
                fail(str(merged), "서로 다른 분석 항목이 병합되어 있습니다. 열을 구분해 주세요.")
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    anchors[(r, c)] = (merged.min_row, merged.min_col)
        if errors:
            raise FmeaImportError(errors)

        def values(col, first, last):
            result, seen = [], set()
            for r in range(first, last + 1):
                anchor = anchors.get((r, col), (r, col))
                cell = sheet.cell(*anchor)
                if anchor in seen or cell.value is None or str(cell.value).strip() == "":
                    continue
                seen.add(anchor)
                if cell.data_type == "f" and col in {12, 20}:
                    # Accept only multiplication of this analysis block's S/O/D cells.
                    formula = compact(cell.value).upper().replace("$", "")
                    body = formula[1:]
                    refs = body[8:-1].split(",") if body.startswith("PRODUCT(") and body.endswith(")") else body.split("*")
                    matches = [re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", ref) for ref in refs]
                    expected = {"D", "H", "K"} if col == 12 else {"Q", "R", "S"}
                    if len(matches) != 3 or not all(matches) or {m[1] for m in matches} != expected or any(not first <= int(m[2]) <= last for m in matches):
                        fail(cell.coordinate, "RPN 수식은 같은 분석행의 S×O×D만 사용할 수 있습니다.")
                        continue
                    factors = [sheet[ref].value for ref in refs]
                    if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not 1 <= v <= 10 or int(v) != v for v in factors):
                        fail(cell.coordinate, "RPN 수식이 참조하는 점수는 1~10의 정수여야 합니다.")
                        continue
                    result.append(int(factors[0] * factors[1] * factors[2]))
                    continue
                if cell.data_type in {"f", "e"}:
                    fail(cell.coordinate, "수식 또는 엑셀 오류값은 이 항목에 사용할 수 없습니다. 계산된 값으로 입력해 주세요.")
                    continue
                result.append(cell.value)
            return result

        def text(col, first, last, required=False, limit=4000):
            result = "\n".join(str(x).strip() for x in values(col, first, last))
            if required and not result:
                fail(f"{get_column_letter(col)}{first}", "필수 분석 내용이 비어 있습니다.")
            if len(result) > limit:
                fail(f"{get_column_letter(col)}{first}", f"입력 내용은 {limit}자 이하여야 합니다.")
            return result

        def number(col, first, last, score=True, required=True):
            data = values(col, first, last)
            if not data and not required:
                return None
            valid = []
            for value in data:
                if isinstance(value, bool) or not re.fullmatch(r"\d+(?:\.0+)?", str(value).strip()):
                    valid = []
                    break
                whole = str(value).strip().split(".")[0].lstrip("0") or "0"
                if len(whole) > 4:
                    valid = []
                    break
                valid.append(int(whole))
            if not valid or len(set(valid)) != 1 or not 1 <= valid[0] <= (10 if score else 1000):
                fail(f"{get_column_letter(col)}{first}", "점수는 1~10의 정수 하나로 입력해 주세요." if score else "RPN은 1~1000의 정수 하나로 입력해 주세요.")
                return None
            return valid[0]

        def target_date(first, last):
            data = values(15, first, last)
            if not data:
                return None
            if len(data) == 1:
                value = data[0]
                if isinstance(value, datetime):
                    return value.date().isoformat()
                if isinstance(value, date):
                    return value.isoformat()
                if isinstance(value, str) and re.fullmatch(r"\d{4}[-./]\d{1,2}[-./]\d{1,2}", value.strip()):
                    try:
                        return date(*map(int, re.split(r"[-./]", value.strip()))).isoformat()
                    except ValueError:
                        pass
            fail(f"O{first}", "완료예정일은 연·월·일을 입력해 주세요(예: 2021-11-30). 월만 있는 날짜는 등록할 수 없습니다.")
            return None

        by_number = {}
        for step in steps:
            key = compact(step.step_no)
            if key in by_number:
                fail("", f"선택한 공정흐름도의 공정번호 {step.step_no}가 중복되어 연결할 수 없습니다.")
            by_number[key] = step
        groups, block_starts = [], []
        for r in range(start, end + 1):
            if any(sheet.cell(r, c).value is not None for c in (7, 14)):
                fail(f"G{r}/N{r}", "병합용 열에 별도 데이터가 있습니다. 회사 양식의 열 배치를 확인해 주세요.")
            if any(sheet.cell(r, c).value is not None for c in range(21, sheet.max_column + 1)):
                fail(f"U{r}", "지원하는 분석 항목 밖에 데이터가 있습니다.")
            a = sheet.cell(r, 1)
            if a.value is not None:
                if a.data_type in {"f", "e"}:
                    fail(a.coordinate, "공정번호·명칭은 수식 없이 입력해 주세요.")
                match = re.fullmatch(r"\s*(\d+)\s+(.+?)\s*", str(a.value), re.S)
                if not match:
                    fail(a.coordinate, "공정번호와 공정명을 구분할 수 없습니다(예: 40 CNC 가공).")
                else:
                    number_key, name = match.groups()
                    step = by_number.get(number_key)
                    if not step:
                        fail(a.coordinate, f"공정번호 {number_key}가 선택한 공정흐름도에 없습니다.")
                    elif compact(name) != compact(step.step_name):
                        fail(a.coordinate, f"공정명 불일치: 엑셀 '{' '.join(name.split())}' / 공정흐름도 '{step.step_name}'.")
                    groups.append((r, number_key, name, step))
            if sheet.cell(r, 2).value is not None:
                block_starts.append(r)
        expected_order = [compact(x.step_no) for x in steps]
        actual_order = [x[1] for x in groups]
        if actual_order != expected_order:
            fail("A", "공정번호·순서·개수가 선택한 공정흐름도와 다릅니다. "
                 f"엑셀: {', '.join(actual_order) or '없음'} / 공정흐름도: {', '.join(expected_order) or '없음'}")
        if not block_starts or len(block_starts) > MAX_ANALYSES:
            fail("B", "분석행은 1~500개여야 합니다.")
        covered = set()
        rows = []
        source_columns = (1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 12, 13, 15, 16, 17, 18, 19, 20)
        for index, first in enumerate(block_starts):
            last = block_starts[index + 1] - 1 if index + 1 < len(block_starts) else end
            b_anchor = anchors.get((first, 2), (first, 2))
            b_last = max((r for (r, c), a in anchors.items() if c == 2 and a == b_anchor), default=first)
            # A block extends through its merged failure-mode cell, not arbitrary following data.
            if last != b_last:
                if any(sheet.cell(r, c).value is not None for r in range(b_last + 1, last + 1) for c in source_columns):
                    fail(f"B{first}", "고장형태가 없는 분석행이 있습니다. 고장형태별 병합 범위를 확인해 주세요.")
                last = b_last
            if b_last >= (block_starts[index + 1] if index + 1 < len(block_starts) else end + 1):
                fail(f"B{first}", "고장형태 병합 범위가 다음 분석행을 침범합니다.")
            covered.update(range(first, last + 1))
            a_anchor = anchors.get((first, 1), (first, 1))
            group = next((g for g in groups if g[0] == a_anchor[0]), None)
            if not group or not group[3] or any(anchors.get((r, 1), (r, 1)) != a_anchor for r in range(first, last + 1)):
                fail(f"A{first}", "분석행의 공정 연결이 비어 있거나 다른 공정의 범위를 넘어갑니다.")
                continue
            step = group[3]
            row = {"flow_step_id": step.id, "function_text": " ".join(group[2].split()),
                   "failure_mode": text(2, first, last, required=True), "effects": text(3, first, last, required=True),
                   "classification": text(5, first, last, limit=50), "causes": text(6, first, last, required=True),
                   "prevention_controls": text(9, first, last), "detection_controls": text(10, first, last),
                   "severity": number(4, first, last), "occurrence": number(8, first, last),
                   "detection": number(11, first, last)}
            original_rpn = number(12, first, last, score=False)
            if all(row[x] is not None for x in ("severity", "occurrence", "detection")):
                product = row["severity"] * row["occurrence"] * row["detection"]
                if original_rpn is not None and original_rpn != product:
                    fail(f"L{first}", f"RPN 불일치: 엑셀 {original_rpn} / S×O×D {product}.")
            action_columns = (13, 15, 16, 17, 18, 19, 20)
            action_values = [v for col in action_columns for v in values(col, first, last)]
            is_na = lambda v: compact(v).upper() == "N/A"
            row["action_not_applicable"] = bool(action_values) and all(is_na(v) for v in action_values)
            if not row["action_not_applicable"]:
                if any(is_na(v) for v in action_values):
                    fail(f"M{first}:T{last}", "조치 내용과 N/A가 혼재되어 있습니다. 조치 해당없음 또는 실제 조치 내용을 일관되게 입력해 주세요.")
                else:
                    row.update(recommended_actions=text(13, first, last), target_date=target_date(first, last),
                               actions_taken=text(16, first, last), new_severity=number(17, first, last, required=False),
                               new_occurrence=number(18, first, last, required=False), new_detection=number(19, first, last, required=False))
                    new_scores = [row[x] for x in ("new_severity", "new_occurrence", "new_detection")]
                    new_rpn = number(20, first, last, score=False, required=False)
                    if any(x is not None for x in new_scores):
                        if any(x is None for x in new_scores):
                            fail(f"Q{first}:S{last}", "조치 후 S/O/D는 세 점수를 모두 입력해 주세요.")
                        elif new_rpn is None or new_rpn != new_scores[0] * new_scores[1] * new_scores[2]:
                            fail(f"T{first}", "조치 후 RPN이 S×O×D와 일치하지 않습니다.")
                    elif new_rpn is not None:
                        fail(f"T{first}", "조치 후 점수 없이 RPN만 입력되어 있습니다.")
            rows.append(row)
        for r in range(start, end + 1):
            if r not in covered and any(sheet.cell(r, c).value is not None for c in source_columns):
                fail(f"B{r}", "고장형태와 연결되지 않은 분석 내용이 있습니다.")
        if errors:
            raise FmeaImportError(errors[:100])
        return rows
    finally:
        workbook.close()
