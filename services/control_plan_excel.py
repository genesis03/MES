"""Import printed lower table only; no macros, external links or arbitrary evaluation."""
import io
import re
import zipfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_BYTES = 10 * 1024 * 1024
TEXT_COLUMNS = {5:'process_detail',6:'equipment',7:'item_no',8:'product',9:'process',10:'classification',
                13:'specification',14:'method',15:'sample_size',16:'sample_frequency',17:'control_method',
                23:'reaction',26:'note'}
BOOL_COLUMNS = {11:'fool_proof',12:'automatic',19:'material',20:'production',21:'quality',22:'engineering'}
HEADERS = {'A11':'공정번호','B11':'공정흐름도','E11':'공정명','F11':'설비명','G11':'관리항목',
           'J11':'특별특성','K11':'대상여부','M11':'관리기준','S11':'관리분담','W11':'이상발생시조치사항',
           'Z11':'비고','B12':'SUB','C12':'MAIN','D12':'외주','G12':'NO','H12':'제품','I12':'공정',
           'K12':'F/P','L12':'자동검사','M12':'규격','N12':'확인방법','O12':'샘플','Q12':'관리방안',
           'S12':'자재','T12':'생산','U12':'QC','V12':'기술','O13':'크기','P13':'주기'}

class ControlPlanImportError(ValueError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__('관리계획서 양식 또는 내용이 일치하지 않습니다.')

def parse_control_plan(data, steps):
    errors = []
    def fail(cell, message):
        errors.append({'cell':cell, 'message':message})
    if len(data) > MAX_BYTES:
        raise ControlPlanImportError([{'cell':'파일','message':'파일 크기는 10MB 이하이어야 합니다.'}])
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 32 * 1024 * 1024:
                raise ValueError('expanded file too large')
        book = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    except Exception as exc:
        raise ControlPlanImportError([{'cell':'파일','message':'읽을 수 있는 .xlsx 파일이 아닙니다.'}]) from exc
    try:
        sheet = book.worksheets[0]
        if sheet.max_row > 2000 or sheet.max_column > 100:
            raise ControlPlanImportError([{'cell':'파일','message':'허용된 양식 범위를 초과했습니다.'}])
        normalize = lambda v: re.sub(r'\s+', '', str(v or ''))
        for cell, label in HEADERS.items():
            if normalize(sheet[cell].value) != normalize(label):
                fail(cell, f'항목명 불일치: {label}')
        if errors:
            raise ControlPlanImportError(errors)
        anchors = {}
        allowed_horizontal = [{2,3,4},{17,18},{23,24,25},{26,27}]
        for merge in sheet.merged_cells.ranges:
            if merge.max_row < 14 or merge.min_col > 27 or sheet.cell(merge.min_row,merge.min_col).value is None:
                continue
            if merge.min_row < 14:
                fail(str(merge), '항목명과 입력 내용이 병합되어 있습니다.')
            if merge.max_col > merge.min_col and not any(set(range(merge.min_col,merge.max_col+1)) <= group for group in allowed_horizontal):
                fail(str(merge), '서로 다른 입력 항목이 병합되어 있습니다.')
            for row in range(merge.min_row, merge.max_row+1):
                for col in range(merge.min_col, min(merge.max_col,27)+1):
                    anchors[(row,col)] = (merge.min_row, merge.min_col)
        cache = {}
        def evaluate(cell, trail=()):
            if cell in cache:
                return cache[cell]
            if cell in trail or len(trail)>30:
                raise ValueError('순환 또는 너무 깊은 수식 참조')
            value = sheet[cell].value
            if sheet[cell].data_type == 'e':
                raise ValueError('엑셀 오류 값')
            if isinstance(value,str) and value.startswith('='):
                formula = value[1:].replace('$','').strip()
                match = re.fullmatch(r'([A-Z]{1,2}[1-9][0-9]*)',formula)
                increment = re.fullmatch(r'([A-Z]{1,2}[1-9][0-9]*)\+([0-9]{1,4})',formula)
                conditional = re.fullmatch(r'IF\(([A-Z]{1,2}[1-9][0-9]*)="","",\1\)',formula,re.I)
                if match or conditional:
                    value = evaluate((match or conditional)[1].upper(),trail+(cell,))
                elif increment:
                    base = evaluate(increment[1],trail+(cell,))
                    if isinstance(base,bool) or not isinstance(base,(int,float)):
                        raise ValueError('번호 수식의 기준 값이 숫자가 아닙니다.')
                    value = base + int(increment[2])
                else:
                    raise ValueError('지원하지 않는 수식입니다. 값으로 변환해 주세요.')
            cache[cell] = value
            return value
        def read(row,col):
            row,col = anchors.get((row,col),(row,col))
            cell = f'{get_column_letter(col)}{row}'
            try:
                value=evaluate(cell)
                if value is None: return ''
                value = str(int(value)) if isinstance(value,float) and value.is_integer() else str(value)
                if len(value)>4000:
                    raise ValueError('내용은 4000자 이하이어야 합니다.')
                return value.strip()
            except ValueError as exc:
                fail(cell,str(exc)); return ''
        rows=[]; groups=[]; current=None; identities=None; seen={}
        for r in range(14,sheet.max_row+1):
            if not any(sheet.cell(r,c).value is not None for c in range(1,28)):
                continue
            number=read(r,1); name=read(r,5)
            if not number or not name:
                fail(f'A{r}', '공정번호와 공정명이 필요합니다.'); continue
            base_name=name.splitlines()[0].strip()
            key=(number,base_name)
            if not groups or groups[-1]!=key:
                groups.append(key)
            # A management number can span several independently named characteristics.
            identity=tuple(anchors.get((r,c),(r,c)) for c in (1,5,6,7,8,9))
            # Blank unmerged cells do not create additional items within a merged item.
            identity=tuple(anchor if read(r,c) else None for anchor,c in zip(identity,(1,5,6,7,8,9)))
            if current is None or identities != identity:
                current={'source_cell':f'A{r}', '_group':key, 'flow_step_id':None,
                         **{f:'' for f in TEXT_COLUMNS.values()}, **{f:False for f in BOOL_COLUMNS.values()},
                         'sub':'','main':'','outside':''}
                rows.append(current); identities=identity; seen={}
            for col,field in TEXT_COLUMNS.items():
                value=read(r,col); anchor=anchors.get((r,col),(r,col))
                if value and anchor not in seen.setdefault(field,set()):
                    current[field] += ('\n' if current[field] else '')+value
                    seen[field].add(anchor)
                    if len(current[field])>4000:
                        fail(f'{get_column_letter(col)}{r}', '병합된 내용은 4000자 이하이어야 합니다.')
            for col,field in BOOL_COLUMNS.items():
                value=read(r,col)
                if value not in {'','●','○','✓','✔','Y','1'}:
                    fail(f'{get_column_letter(col)}{r}', '체크 표시는 ●, ○, ✓, ✔, Y, 1 또는 공란을 사용해 주세요.')
                current[field] = current[field] or bool(value)
            for col,field in [(2,'sub'),(3,'main'),(4,'outside')]:
                value=read(r,col)
                if value and current[field] and current[field]!=value:
                    fail(f'{get_column_letter(col)}{r}','같은 관리항목에 공정기호가 다릅니다.')
                if value: current[field]=value
        # Consecutive material-specific blocks are one process in the flow.
        expected=[(str(s.step_no).strip(),str(s.step_name).strip()) for s in steps]
        if len(set(expected))!=len(expected):
            fail('공정흐름도','동일한 공정번호와 공정명이 중복되어 엑셀 행을 구분할 수 없습니다.')
        if groups!=expected:
            fail('A14:E마지막행','공정번호·공정명 또는 순서가 공정흐름도와 일치하지 않습니다. 누락된 공정도 확인해 주세요.')
        mapping={key:s.id for key,s in zip(expected,steps)}
        for row in rows:
            row['flow_step_id']=mapping.get(row.pop('_group'))
            if not (row['product'] or row['process']):
                fail(row['source_cell'],'제품 또는 공정 관리항목이 필요합니다.')
            row.pop('source_cell')
        if not rows or len(rows)>500:
            fail('입력행','관리항목은 1~500행이어야 합니다.')
        if errors:
            raise ControlPlanImportError(errors[:100])
        return rows
    finally:
        book.close()
