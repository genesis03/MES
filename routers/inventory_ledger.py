import json
import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, Request, Query
from fastapi.responses import StreamingResponse
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field, field_validator, model_validator
from sqlalchemy import update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from core.database import get_db
from core import security
from core.security import get_current_user, check_admin_permission, parse_user_permissions
from models.models import ItemMasterModel
from models.inventory_ledger import InventoryLedgerBalance, InventoryLedgerEntry, InventoryLedgerAudit
from services.inventory_ledger_service import report,detail,month_bounds

router=APIRouter(tags=['Inventory Ledger'])
templates=Jinja2Templates(directory='templates')
MENU='/inventory/ledger'
Source=Literal['mes','external']


def permission(user,write=False):
    if security.DEV_BYPASS_AUTH or check_admin_permission(user):return
    if not user:raise HTTPException(401,'로그인이 필요합니다.')
    access=parse_user_permissions(user).get('menu_access',{})
    level=access.get(MENU,'NONE')
    allowed=level in ('WRITE',) if write else level in ('READ','WRITE',True)
    if not allowed:raise HTTPException(403,'재고 수불·실사 '+('수정' if write else '조회')+' 권한이 없습니다.')


def valid_month(value):
    if not re.fullmatch(r'\d{4}-\d{2}',value):raise ValueError('월은 YYYY-MM 형식이어야 합니다.')
    month_bounds(value)
    return value


class Filter(BaseModel):
    month:str
    source:Source='mes'
    @field_validator('month')
    @classmethod
    def month_format(cls,v):return valid_month(v)


def ledger_filters(month:str=Query(...,pattern=r'^\d{4}-\d{2}$'),source:Source=Query('mes')):
    try:return Filter(month=month,source=source)
    except ValueError:raise HTTPException(422,'유효한 조회 월을 선택하세요.')


class BalanceInput(Filter):
    item_id:int=Field(gt=0)
    opening_qty:Decimal|None=Field(default=None,ge=0,max_digits=20,decimal_places=6,allow_inf_nan=False)
    actual_qty:Decimal|None=Field(default=None,ge=0,max_digits=20,decimal_places=6,allow_inf_nan=False)
    count_carry_qty:Decimal=Field(default=Decimal(0),ge=0,max_digits=20,decimal_places=6,allow_inf_nan=False)
    count_date:date|None=None
    note:str=Field(min_length=1,max_length=1000)
    version:int=Field(ge=0)
    @model_validator(mode='after')
    def check_count(self):
        if self.count_carry_qty and self.actual_qty is None:raise ValueError('이월수량에는 실사수량이 필요합니다.')
        if self.actual_qty is not None and self.count_date is None:raise ValueError('실사수량에는 실사일이 필요합니다.')
        if self.count_date and self.count_date.strftime('%Y-%m')!=self.month:raise ValueError('실사일은 조회 월 안의 날짜여야 합니다.')
        if not self.note.strip():raise ValueError('변경 사유를 입력하세요.')
        return self


class EntryInput(BaseModel):
    source:Source
    entry_date:date
    item_id:int=Field(gt=0)
    direction:Literal['IN','OUT']
    category:Literal['구매입고','외주출고','외주입고','공정투입','포장투입','외주불량','LOSS','샘플','킵업','기타']
    qty:Decimal=Field(gt=0,max_digits=20,decimal_places=6,allow_inf_nan=False)
    document_no:str=Field(default='',max_length=100)
    lot_no:str=Field(default='',max_length=100)
    reason:str=Field(min_length=1,max_length=1000)
    @model_validator(mode='after')
    def correct_direction(self):
        expected = 'IN' if self.category in ('구매입고','외주입고') else 'OUT'
        if self.category != '기타' and self.direction != expected:
            raise ValueError('선택한 수불 유형의 입출고 구분을 확인하세요.')
        return self
    @field_validator('reason')
    @classmethod
    def reason_nonblank(cls,v):
        if not v.strip():raise ValueError('보완 사유를 입력하세요.')
        return v.strip()


class VoidInput(BaseModel):
    reason:str=Field(min_length=1,max_length=1000)
    @field_validator('reason')
    @classmethod
    def nonblank(cls,v):
        if not v.strip():raise ValueError('취소 사유를 입력하세요.')
        return v.strip()


def item_required(db,item_id):
    item=db.get(ItemMasterModel,item_id)
    if not item:raise HTTPException(404,'품번을 찾을 수 없습니다.')
    return item


def audit(db,month,source,item_id,action,before,after,user):
    db.add(InventoryLedgerAudit(month=month,source=source,item_id=item_id,action=action,
        before_json=json.dumps(before,ensure_ascii=False,default=str),after_json=json.dumps(after,ensure_ascii=False,default=str),
        created_by=user.username))


@router.get(MENU)
def page(request:Request,user=Depends(get_current_user)):
    permission(user)
    try:permission(user,True);can_write=True
    except HTTPException:can_write=False
    return templates.TemplateResponse(request=request,name='inventory_ledger.html',context={'request':request,'user':user,'can_write':can_write})


@router.get('/api/inventory/ledger')
def get_report(filters:Filter=Depends(ledger_filters),part_no:str=Query('',max_length=100),category:str=Query('',max_length=50),db:Session=Depends(get_db),user=Depends(get_current_user)):
    permission(user)
    return report(db,filters.month,filters.source,part_no,category)[0]


@router.get('/api/inventory/ledger/detail')
def get_detail(filters:Filter=Depends(ledger_filters),item_id:int|None=Query(None,gt=0),raw_part:str=Query('',max_length=200),process:str=Query('',max_length=200),db:Session=Depends(get_db),user=Depends(get_current_user)):
    permission(user)
    result=detail(db,filters.month,filters.source,item_id,raw_part,process)
    if result is None:raise HTTPException(404,'품번을 찾을 수 없습니다.')
    return result


@router.put('/api/inventory/ledger/balance')
def save_balance(payload:BalanceInput,db:Session=Depends(get_db),user=Depends(get_current_user)):
    permission(user,True);item_required(db,payload.item_id)
    row=db.query(InventoryLedgerBalance).filter_by(month=payload.month,source=payload.source,item_id=payload.item_id).one_or_none()
    before={k:getattr(row,k) for k in ('opening_qty','actual_qty','count_carry_qty','count_date','note','version')} if row else {}
    values=payload.model_dump(exclude={'month','source','item_id','version'})
    values['count_date']=payload.count_date.isoformat() if payload.count_date else None
    values.update(version=payload.version+1,updated_at=datetime.now(),updated_by=user.username)
    try:
        if row:
            result=db.execute(update(InventoryLedgerBalance).where(InventoryLedgerBalance.id==row.id,InventoryLedgerBalance.version==payload.version).values(**values))
            if result.rowcount!=1:raise HTTPException(409,'다른 사용자가 수정했습니다. 다시 조회 후 저장하세요.')
        else:
            if payload.version!=0:raise HTTPException(409,'이력이 변경되었습니다. 다시 조회하세요.')
            db.add(InventoryLedgerBalance(month=payload.month,source=payload.source,item_id=payload.item_id,**values))
        audit(db,payload.month,payload.source,payload.item_id,'BALANCE',before,values,user)
        db.commit()
    except IntegrityError:
        db.rollback();raise HTTPException(409,'다른 사용자가 저장했습니다. 다시 조회하세요.')
    return {'status':'success','version':payload.version+1}


@router.post('/api/inventory/ledger/entries')
def add_entry(payload:EntryInput,db:Session=Depends(get_db),user=Depends(get_current_user)):
    permission(user,True);item_required(db,payload.item_id)
    values=payload.model_dump();values['entry_date']=payload.entry_date.isoformat()
    row=InventoryLedgerEntry(**values,created_by=user.username)
    db.add(row);db.flush()
    audit(db,payload.entry_date.strftime('%Y-%m'),payload.source,payload.item_id,'ENTRY',{},values|{'id':row.id},user)
    db.commit()
    return {'status':'success','id':row.id}


@router.post('/api/inventory/ledger/entries/{entry_id}/void')
def void_entry(entry_id:int,payload:VoidInput,db:Session=Depends(get_db),user=Depends(get_current_user)):
    permission(user,True)
    row=db.get(InventoryLedgerEntry,entry_id)
    if not row:raise HTTPException(404,'수불 보완 이력이 없습니다.')
    result=db.execute(update(InventoryLedgerEntry).where(InventoryLedgerEntry.id==entry_id,InventoryLedgerEntry.voided.is_(False)).values(voided=True))
    if result.rowcount:
        audit(db,row.entry_date[:7],row.source,row.item_id,'VOID',{'id':row.id,'qty':row.qty},{'reason':payload.reason},user)
    db.commit();return {'status':'success'}


@router.get('/api/inventory/ledger/history')
def history(filters:Filter=Depends(ledger_filters),item_id:int=Query(...,gt=0),db:Session=Depends(get_db),user=Depends(get_current_user)):
    permission(user)
    rows=db.query(InventoryLedgerAudit).filter_by(month=filters.month,source=filters.source,item_id=item_id).order_by(InventoryLedgerAudit.id.desc()).all()
    return [{'at':r.created_at.isoformat(timespec='seconds'),'by':r.created_by,'action':r.action,'before':json.loads(r.before_json),'after':json.loads(r.after_json)} for r in rows]


@router.get('/api/inventory/ledger/export')
def export(filters:Filter=Depends(ledger_filters),part_no:str=Query('',max_length=100),category:str=Query('',max_length=50),db:Session=Depends(get_db),user=Depends(get_current_user)):
    permission(user)
    from openpyxl import Workbook
    data,events=report(db,filters.month,filters.source,part_no,category)
    book=Workbook();sheet=book.active;sheet.title='월간 수불·실사'
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    columns=[('구분','category_name'),('품번','part_no'),('단위','unit'),('기초','opening_qty'),('입고','in_qty'),
             ('생산','production_qty'),('불량','production_defect_qty'),('Setup','setup_qty'),
             ('입고','subcontract_in_qty'),('불량','subcontract_defect_qty'),('LOSS','loss_qty'),('샘플','sample_qty'),
             ('포장','packing_qty'),('원재료','raw_in_qty'),('출고','out_qty'),('기말(산출)','closing_qty'),
             ('실재고','actual_qty'),('이월수량','count_carry_qty'),('차이 (실사-산출)','difference')]
    for column in [1,2,3,4,5,15,16,19]:
        sheet.cell(1,column,columns[column-1][0]);sheet.merge_cells(start_row=1,start_column=column,end_row=3,end_column=column)
    sheet.merge_cells('F1:N1');sheet.merge_cells('F2:H2');sheet['F2']='생산'
    sheet.merge_cells('I2:L2');sheet['I2']='외주가공'
    for column in (13,14):
        sheet.cell(2,column,columns[column-1][0]);sheet.merge_cells(start_row=2,start_column=column,end_row=3,end_column=column)
    sheet.merge_cells('Q1:R1');sheet['Q1']='재고실사'
    for column in (17,18):
        sheet.cell(2,column,columns[column-1][0]);sheet.merge_cells(start_row=2,start_column=column,end_row=3,end_column=column)
    for column in range(6,13):sheet.cell(3,column,columns[column-1][0])
    def safe(v):return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v
    numeric={key for _,key in columns if key not in ('category_name','part_no','unit')}
    for row in data['rows']:
        sheet.append([safe(row[key]) if key not in numeric else (Decimal(row[key]) if row[key] is not None else None) for _,key in columns])
        sheet.cell(sheet.max_row,2).comment=Comment(f"품명: {row['part_name']}\n실사일: {row['count_date'] or '미입력'}\n실사일 장부재고: {row['expected_at_count']}\n사유: {row['note']}",'MES')
    # Preserve name/date/memo details without adding columns to the requested form.
    survey=book.create_sheet('실사 상세')
    survey_columns=[('품번','part_no'),('품명','part_name'),('단위','unit'),('실사일','count_date'),('실재고','actual_qty'),('이월수량','count_carry_qty'),('실사일 장부재고','expected_at_count'),('차이','difference'),('사유','note')]
    survey.append([label for label,_ in survey_columns])
    for row in data['rows']:
        survey.append([(Decimal(row[key]) if row[key] is not None else None) if key in numeric or key=='expected_at_count' else safe(row[key]) for _,key in survey_columns])
    evidence=book.create_sheet('수불 상세');evidence.append(['일자','품번','공정','구분','입출고','수량','전표','LOT','비고'])
    ids={r['item_id'] for r in data['rows'] if r['linked']};rawkeys={(r['part_no'],r['process']) for r in data['rows'] if not r['linked']}
    master={r['item_id']:r['part_no'] for r in data['rows'] if r['linked']}
    for e in events:
        if e['item_id'] not in ids and (e['raw_part'],e['process']) not in rawkeys:continue
        evidence.append([safe(v) for v in [e['date'],master.get(e['item_id'],e['raw_part']),e['process'],e['category'],e['direction'],e['qty'],e['document_no'],e['lot_no'],e['note']]])
    daily=book.create_sheet('일별 수불');daily.append(['일자','품번','품명','단위','입고','출고','계산재고'])
    from datetime import timedelta
    grouped={}
    for e in events:
        key=e['item_id'] if e['item_id'] is not None else (e['raw_part'],e['process'])
        values=grouped.setdefault((key,e['date']),{'IN':Decimal(0),'OUT':Decimal(0)})
        values[e['direction']]+=e['qty']
    for row in data['rows']:
        key=row['item_id'] if row['linked'] else (row['part_no'],row['process'])
        running=Decimal(row['opening_qty']) if row['opening_qty'] is not None else None
        day=date.fromisoformat(data['first'])
        while day.isoformat()<=data['last']:
            amounts=grouped.get((key,day.isoformat()),{'IN':Decimal(0),'OUT':Decimal(0)})
            if running is not None:running+=amounts['IN']-amounts['OUT']
            daily.append([day.isoformat(),safe(row['part_no']),safe(row['part_name']),safe(row['unit']),amounts['IN'],amounts['OUT'],running])
            day+=timedelta(days=1)
    types=book.create_sheet('유형별 집계')
    categories=sorted({key for row in data['rows'] for key in row['totals']})
    types.append(['품번','품명','단위']+categories)
    for row in data['rows']:types.append([safe(row['part_no']),safe(row['part_name']),safe(row['unit'])]+[Decimal(row['totals'].get(key,'0')) for key in categories])
    notes=book.create_sheet('집계 기준');notes.append(['자료 구분',filters.source]);notes.append(['월',filters.month])
    for warning in data['warnings']:notes.append([warning])
    for ws in book:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for column in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(column)].width=20
    sheet.freeze_panes='E4';sheet.auto_filter.ref=None
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment=Alignment(horizontal='center' if cell.row<=3 else 'right',vertical='center',wrap_text=cell.row<=3)
            cell.border=Border(left=Side(style='thin',color='94A3B8'),right=Side(style='thin',color='94A3B8'),top=Side(style='thin',color='94A3B8'),bottom=Side(style='thin',color='94A3B8'))
            if cell.row<=3:cell.fill=PatternFill('solid',fgColor='D9D9D9' if cell.row==2 and 6<=cell.column<=14 else 'BDDFE7')
    for index in (1,2,3):sheet.row_dimensions[index].height=22
    for column in range(1,20):sheet.column_dimensions[get_column_letter(column)].width=12
    sheet.column_dimensions['B'].width=22;sheet.column_dimensions['S'].width=18
    stream=BytesIO();book.save(stream);stream.seek(0)
    return StreamingResponse(stream,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="inventory-ledger-{filters.month}-{filters.source}.xlsx"'})
