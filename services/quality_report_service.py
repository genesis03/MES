"""Bind inspection snapshots to the configured workbook without rebuilding styles."""
from io import BytesIO
from pathlib import Path
from statistics import mean
from openpyxl import load_workbook
from core.config import BASE_DIR

TEMPLATE_PATH = BASE_DIR / 'assets' / 'InspectionReport_Template.xlsx'


def evaluate(values, lower, upper, manual=''):
    present = [v for v in values if v is not None]
    results = [None if v is None or (lower is None and upper is None) else
               ('NG' if (lower is not None and v < lower) or (upper is not None and v > upper) else 'OK')
               for v in values]
    average = mean(present) if present else None
    if 'NG' in results:
        judgment = 'NG'
    elif lower is None and upper is None:
        judgment = manual if manual in {'OK', 'NG'} else 'HOLD'
    else:
        judgment = 'OK' if len(present) == 5 else 'HOLD'
    return average, judgment, results


def report_dict(report):
    fields = ('id', 'report_no', 'report_type', 'item_id', 'part_no', 'part_name', 'lot_no', 'order_no',
              'inspection_date', 'inspector_name', 'standard_id', 'standard_revision', 'sample_qty', 'judgment', 'note')
    result = {key: getattr(report, key) for key in fields}
    result['items'] = []
    for item in report.items:
        row = {key: getattr(item, key) for key in ('item_seq', 'item_name', 'item_category', 'spec_display',
               'spec_min', 'spec_max', 'eval_tool', 'avg_val', 'judgment')}
        row['values'] = [getattr(item, f'val_{i}') for i in range(1, 6)]
        row['sample_judgments'] = evaluate(row['values'], item.spec_min, item.spec_max)[2]
        result['items'].append(row)
    return result


def excel_text(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'


def build_report_excel(report, template_path=None):
    workbook = load_workbook(Path(template_path or TEMPLATE_PATH))
    sheet = workbook.active
    data = report_dict(report)
    for cell, key in {'C3': 'report_no', 'C4': 'part_no', 'C5': 'part_name', 'G3': 'inspection_date',
                      'G4': 'lot_no', 'G5': 'sample_qty', 'K4': 'inspector_name', 'C39': 'note'}.items():
        excel_text(sheet[cell], data[key] or '')
    for row_index in range(10, 37):
        for col in range(2, 14):
            sheet.cell(row_index, col).value = None
    for index, item in enumerate(data['items'], 10):
        values = [item['item_seq'], item['item_category'], item['item_name'], item['spec_display'],
                  item['eval_tool'], *item['values'], item['avg_val'], item['judgment']]
        for column, value in enumerate(values, 2):
            excel_text(sheet.cell(index, column), value)
    excel_text(sheet['L38'], {'PASS': '합격', 'REJECT': '불합격', 'HOLD': '보류'}[report.judgment])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
