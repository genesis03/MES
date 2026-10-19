from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_plan_excel(data, start, end):
    book = Workbook()
    summary = book.active
    summary.title = '공정별 요약'
    processes = data['processes']
    summary.append(['생산계획', f'{start} ~ {end}'])
    summary.append(['차종', '품번', '품명', *[p['name'] for p in processes], '합계'])
    for row in data['summary']:
        summary.append([row['vehicle_model'], row['part_no'], row['part_name'],
                        *[row['by_process'].get(p['code'], 0) for p in processes], row['total']])
    summary.append(['합계', '', '', *[sum(r['by_process'].get(p['code'], 0) for r in data['summary']) for p in processes], data['total']])
    detail = book.create_sheet('일자별 생산계획')
    detail.append(['생산계획', f'{start} ~ {end}', '저장된 계획 기준'])
    detail.append(['설비', '공정', '품번', '품명', 'C/T(초)', '목표 OEE', 'UPH', 'UPD(12H)',
                   '기간 생산성(주간)', '기간 생산성(주/야)', '합계', *data['dates']])
    for row in data['rows']:
        detail.append([row['equipment_name'], row['process_name'], row['part_no'], row['part_name'],
                       row['cycle_seconds'], row['expected_oee'], row['uph'], row['upd'],
                       row['upm_day'], row['upm_full'], row['total'],
                       *[row['quantities'].get(d, 0) for d in data['dates']]])
        detail.cell(detail.max_row, 6).number_format = '0.00%'
    detail.append(['합계', '', '', '', '', '', '', '', '', '', data['total'],
                   *[sum(r['quantities'].get(d, 0) for r in data['rows']) for d in data['dates']]])
    for sheet in book:
        sheet.freeze_panes = 'L3' if sheet == detail else 'D3'
        sheet.auto_filter.ref = f'A2:{get_column_letter(sheet.max_column)}{max(2, sheet.max_row - 1)}'
        sheet.print_title_rows = '1:2'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if cell.row in (1, 2, sheet.max_row):
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill('solid', fgColor='E7EDF5')
                cell.alignment = Alignment(vertical='center', wrap_text=cell.row == 2)
                if isinstance(cell.value, (float, int)) and cell.number_format == 'General':
                    cell.number_format = '#,##0.##'
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 16 if col < 12 else 12
        sheet.column_dimensions['D' if sheet == detail else 'C'].width = 40
    output = BytesIO()
    book.save(output)
    return output.getvalue()
