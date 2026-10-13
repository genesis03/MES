import io
import json
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook
from test_production_sync import setup,DAY
from test_production_daily_sync import daily_response
from services.production_daily_sync import parse_daily,MACHINES,SHIFTS,DATA
from services import production_sync_client as source,production_sync_service as sync
from models.production_sync import ProductionSyncFailureData,ExternalProductionRecord,ProductionSyncRun
from core.security import get_current_user


def test_failed_query_snapshot_and_excel_show_original_blank_lot_without_credentials(setup,monkeypatch):
    values=daily_response()
    values[DATA+'LOT_NUM']=['']
    values[DATA+'PRODUCT_NM']=['=1+1']
    values['Common.Set.LoginUserSet.USER_PASS']=['private-password']
    values[DATA+'USER_PASS']=['private-password']
    values['THROW.STACK']=['private SQL and cookie']
    class Client:
        def __init__(self,*_):pass
        def fetch_day(self,day):return parse_daily(values,day,MACHINES[0],SHIFTS[0])
    monkeypatch.setattr(sync,'ProductionClient',Client)
    token,run_id=sync.claim(DAY,DAY,'MANUAL');sync.execute(DAY,DAY,token,run_id)
    with setup.sessions() as db:
        assert db.get(ProductionSyncRun,run_id).status=='FAILED'
        snapshot=db.get(ProductionSyncFailureData,run_id)
        assert snapshot is not None and 'private' not in snapshot.data_json
        data=json.loads(snapshot.data_json)
        assert data['equipment']=='LT-01' and data['fields'][DATA+'LOT_NUM']==['']
        assert db.query(ExternalProductionRecord).count()==0
    runs=setup.client.get('/api/production/external-sync/runs').json()
    assert runs[0]['has_error_data'] is True
    response=setup.client.get(f'/api/production/external-sync/runs/{run_id}/error-excel')
    assert response.status_code==200 and 'attachment' in response.headers['content-disposition']
    workbook=load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames==['조회 및 오류','원본 항목','쉼표 분리 참고']
    rows=list(workbook['원본 항목'].iter_rows(min_row=2,values_only=True))
    lot=next(r for r in rows if r[0]==DATA+'LOT_NUM')
    assert lot[1]=='공란' and lot[3]==0
    product=next(cells[5] for cells in workbook['원본 항목'].iter_rows(min_row=2) if cells[0].value==DATA+'PRODUCT_NM')
    assert product.value=='=1+1' and product.data_type=='s'
    assert all('private' not in str(cell.value) for sheet in workbook for cells in sheet for cell in cells)
    setup.app.dependency_overrides[get_current_user]=lambda:SimpleNamespace(username='reader',role='USER',permissions='{}')
    assert setup.client.get(f'/api/production/external-sync/runs/{run_id}/error-excel').status_code==403


def test_excel_preserves_long_comma_note_chunks_and_does_not_claim_row_alignment():
    from services.production_sync_error_excel import build_error_excel
    note=('raw note, with punctuation\n'*1800)
    data={'work_date':DAY.isoformat(),'equipment':'LT-01','error':'bad count','fields':{DATA+'PRODUCT_NM':[note]}}
    workbook=load_workbook(io.BytesIO(build_error_excel(data)))
    rows=[r for r in workbook['원본 항목'].iter_rows(min_row=2,values_only=True) if r[0]==DATA+'PRODUCT_NM']
    assert ''.join(r[5] or '' for r in rows)==note
    assert len(rows)>1
    missing=next(r for r in workbook['원본 항목'].iter_rows(min_row=2,values_only=True) if r[0]==DATA+'LOT_NUM')
    assert missing[1]=='미수신'
    assert '행 연결 미확정' in workbook['쉼표 분리 참고']['E1'].value


def test_existing_history_without_snapshot_has_no_download(setup):
    token,run_id=sync.claim(DAY,DAY,'MANUAL')
    assert setup.client.get('/api/production/external-sync/runs').json()[0]['has_error_data'] is False
    response=setup.client.get(f'/api/production/external-sync/runs/{run_id}/error-excel')
    assert response.status_code==404 and '다시 가져오면' in response.json()['detail']


def test_assembly_review_preserves_comparison_and_is_admin_only(setup,monkeypatch):
    from services.production_daily_sync import supplement_assembly,ASSEMBLY_MACHINES,COND
    from test_production_daily_sync import assembly_lots,assembly_response,daily_response
    class Client:
        def __init__(self,*_):self.opener=object();self.context={}
        def fetch_day(self,day):
            legacy=assembly_lots();legacy[0]['JOB_QTY']='1'
            return legacy,supplement_assembly(self,day,legacy)
    monkeypatch.setattr(source,'query_with_retry',lambda _,payload:assembly_response() if payload[COND+'MACH_NM']=='AS-01' else daily_response(machine=ASSEMBLY_MACHINES[1],empty=True))
    monkeypatch.setattr(sync,'ProductionClient',Client)
    token,run_id=sync.claim(DAY,DAY,'MANUAL');sync.execute(DAY,DAY,token,run_id)
    history=setup.client.get('/api/production/external-sync/runs').json()[0]
    assert len(history['counts']['review_items'])==1
    assert '미연결 1건' in history['counts']['warnings'][0]
    with setup.sessions() as db:
        assert db.get(ProductionSyncRun,run_id).status=='SUCCESS'
        assert db.query(ExternalProductionRecord).count()==2
    path=f'/api/production/external-sync/runs/{run_id}/assembly-review-excel'
    response=setup.client.get(path)
    assert response.status_code==200 and 'attachment' in response.headers['content-disposition']
    workbook=load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames==['조립 확인','LOT 비교']
    assert workbook['조립 확인']['H2'].value=='4560'
    assert workbook['조립 확인']['K2'].value==1441
    assert workbook['LOT 비교'].max_row==3
    setup.app.dependency_overrides[get_current_user]=lambda:SimpleNamespace(username='reader',role='USER',permissions='{}')
    assert setup.client.get(path).status_code==403


def test_assembly_review_excel_does_not_execute_formula_text():
    from services.production_sync_error_excel import build_assembly_review_excel
    review=dict(work_date='2026-03-21',equipment='AS-01',shift='주간',part_no='310061',reason='수량 불일치',daily=dict(OPERATOR_NAME='=1+1',JOB_QTY='10',FAULT_QTY='5',F10='5',JOB_ST_TIME='',JOB_END_TIME=''),lots=[])
    workbook=load_workbook(io.BytesIO(build_assembly_review_excel([review])))
    assert workbook['조립 확인']['G2'].data_type=='s'
    assert workbook['조립 확인']['G2'].value=='=1+1'
