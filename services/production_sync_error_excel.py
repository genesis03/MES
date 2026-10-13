"""Export a captured failing response without inventing row associations or exposing login data."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from services.production_daily_sync import DATA


def build_error_excel(data):
    workbook=Workbook(write_only=True)
    def sheet(name,headers,widths):
        ws=workbook.create_sheet(name);ws.freeze_panes='A2'
        for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
        cells=[]
        for label in headers:
            c=WriteOnlyCell(ws,value=label);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476B');cells.append(c)
        ws.append(cells);return ws
    def append(ws,values):
        cells=[]
        for value in values:
            c=WriteOnlyCell(ws,value=value)
            if isinstance(value,str):c.data_type='s'
            c.alignment=Alignment(vertical='top',wrap_text=True);cells.append(c)
        ws.append(cells)
    info=sheet('조회 및 오류',['항목','내용'],[25,110])
    for label,key in [('작업일','work_date'),('설비','equipment'),('공정','process'),('주야','shift'),('오류','error')]:append(info,[label,data.get(key,'')])
    append(info,['자료 범위','오류가 발생한 한 조회 조건의 원본 응답입니다. 정상 저장된 전체 생산실적 목록이 아닙니다.'])
    append(info,['확인 방법','원본 항목에서 LOT·품번·수량·시간을 확인하세요. 쉼표는 행 구분자와 문장 부호로 함께 사용되므로 분리 수가 실제 건수를 뜻하지 않을 수 있습니다.'])
    append(info,['긴 값','Excel 셀 길이 제한을 넘는 원문은 여러 조각으로 나눠 저장합니다. 조각 순서대로 연결하면 원문이 됩니다.'])
    raw=sheet('원본 항목',['항목','응답 상태','반환값 번호','쉼표 분리 수(참고)','원문 조각 번호','원문 값'],[75,15,15,25,17,100])
    pieces=sheet('쉼표 분리 참고',['항목','반환값 번호','분리 순번','값 조각 번호','분리 값 (행 연결 미확정)'],[75,15,15,17,85])
    fields=data.get('fields',{})
    required={DATA+k for k in 'ITEM_NUM LOT_NUM RULE_NO JOB_PROD_QTY JOB_FAULT_QTY SET_UP_QTY JOB_TIME_STRING'.split()}
    for key in sorted(set(fields)|required):
        if key not in fields:append(raw,[key,'미수신',None,None,None,None]);continue
        for number,value in enumerate(fields[key],1):
            tokens=value.split(',') if value else []
            chunks=[value[i:i+30000] for i in range(0,len(value),30000)] or ['']
            for chunk_number,chunk in enumerate(chunks,1):append(raw,[key,'값 있음' if value else '공란',number,len(tokens),chunk_number,chunk])
            for index,token in enumerate(tokens,1):
                for chunk_index,start in enumerate(range(0,max(1,len(token)),30000),1):append(pieces,[key,number,index,chunk_index,token[start:start+30000]])
    buffer=BytesIO();workbook.save(buffer);return buffer.getvalue()


def build_assembly_review_excel(reviews):
    """Compare observed assembly totals with candidate LOTs; never infer allocations."""
    workbook=Workbook()
    workbook.remove(workbook.active)
    summary=workbook.create_sheet('조립 확인')
    lots=workbook.create_sheet('LOT 비교')
    headers=['번호','작업일','설비','주야','품번','연결 확인사유','작업자','작업일보 양품','작업일보 불량','작업일보 셋업','LOT 합계 양품','LOT 합계 불량','LOT 합계 셋업','작업일보 시작','작업일보 종료']
    def append(ws,values):
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    append(summary,headers)
    append(lots,['확인 번호','작업일','설비','품번','LOT 번호','작업번호','양품','불량','셋업','시작','종료'])
    for index,review in enumerate(reviews,1):
        daily=review['daily'];candidates=review['lots']
        totals=[sum(float(r[k]) for r in candidates) for k in ('JOB_QTY','FAULT_QTY','F10')]
        append(summary,[index,review['work_date'],review['equipment'],review['shift'],review['part_no'],review['reason'],daily['OPERATOR_NAME'],
                        daily['JOB_QTY'],daily['FAULT_QTY'],daily['F10'],*totals,daily['JOB_ST_TIME'],daily['JOB_END_TIME']])
        for row in candidates:
            append(lots,[index,review['work_date'],review['equipment'],review['part_no'],row['LOT_NUM'],row['JOB_NUM'],row['JOB_QTY'],row['FAULT_QTY'],row['F10'],row['JOB_ST_TIME'],row['JOB_END_TIME']])
    for ws in workbook:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476B')
        for column in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(column)].width=24 if column>1 else 10
    buffer=BytesIO();workbook.save(buffer);return buffer.getvalue()
