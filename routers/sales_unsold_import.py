"""Reviewed replacement of outstanding orders; shipment history is retained."""
import hashlib
import io
import json
import math
import re
from datetime import datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, Request
from openpyxl import load_workbook
from fastapi.templating import Jinja2Templates
from sqlalchemy import text
from sqlalchemy.orm import Session

from core.database import get_db
from core.security import require_admin_user
from models.models import ItemMasterModel
from models.partner import Partner
from models.sales import SalesOrderMaster, SalesOrderItem
from routers.sales_shipping_entry import _next_no, _username

router = APIRouter(tags=['Sales outstanding import'])


def name_key(value):
    return re.sub(r'\s+', '', str(value or '')).replace('(주)', '').replace('㈜', '')


def read_rows(data):
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheet = workbook.active
        if sheet.max_row > 10000 or sheet.max_column > 100:
            raise ValueError('엑셀 크기가 너무 큽니다.')
        values = list(sheet.values)
        workbook.close()
        if tuple(values[1][:6]) != ('일자-No.', '거래처명', '품번', '품목명', '수주수량', '미판매수량'):
            raise ValueError('이카운트 미판매현황 엑셀을 선택해 주세요.')
        rows = []
        keys = set()
        for number, row in enumerate(values[2:], 3):
            if not re.match(r'^\d{4}/\d{2}/\d{2}\s*-', str(row[0])):
                if row[2] is not None or row[3] is not None:
                    raise ValueError(f'{number}행: 수주일자를 확인해 주세요.')
                continue
            order_date = datetime.strptime(str(row[0]).split(' -')[0].strip(), '%Y/%m/%d').strftime('%Y-%m-%d')
            qty = float(row[5])
            if not math.isfinite(qty) or qty <= 0:
                raise ValueError(f'{number}행: 미판매수량은 양수여야 합니다.')
            due = str(row[9] or '').strip()
            if due:
                due = datetime.strptime(due, '%Y/%m/%d').strftime('%Y-%m-%d')
            key = (str(row[0]).strip(), str(row[1]).strip(), str(row[2]).strip())
            if key in keys:
                raise ValueError(f'{number}행: 동일 수주·거래처·품번이 중복되었습니다.')
            keys.add(key)
            rows.append(dict(source=key[0], customer=key[1], part_no=key[2], order_date=order_date,
                             qty=qty, due=due or None, note=str(row[8] or '')))
        if not rows:
            raise ValueError('등록할 상세 항목이 없습니다.')
        totals = [float(r[5]) for r in values if str(r[0]).strip() == '총합계']
        if len(totals) != 1 or abs(totals[0] - sum(r['qty'] for r in rows)) > 1e-6:
            raise ValueError('상세 수량과 엑셀 총합계가 일치하지 않습니다.')
        return rows
    except Exception as exc:
        raise HTTPException(422, f'엑셀 확인 실패: {exc}') from exc


def completed(db):
    return bool(db.query(SalesOrderMaster.id).filter(SalesOrderMaster.po_no.like('UNSOLD-XLSX:%')).first())


@router.get('/admin/sales-unsold-migration')
def migration_page(request: Request, db: Session = Depends(get_db), user=Depends(require_admin_user)):
    return Jinja2Templates(directory='templates').TemplateResponse(request=request, name='sales_unsold_migration.html',
        context={'request': request, 'user': user, 'completed': completed(db)})


def prepare(db, data):
    if completed(db):
        raise HTTPException(409, '미판매 잔량 일회성 전환이 이미 완료되었습니다. 재실행할 수 없습니다.')
    rows = read_rows(data)
    customers = {}
    for partner in db.query(Partner).filter(Partner.is_active == 'Y', Partner.partner_type.in_(['CUSTOMER', 'BOTH'])):
        customers.setdefault(name_key(partner.partner_name), []).append(partner)
    items = {i.part_no: i for i in db.query(ItemMasterModel).filter(ItemMasterModel.is_active == 'Y', ItemMasterModel.material_type.in_(['SEMI', 'FINISHED']))}
    for row in rows:
        matches = customers.get(name_key(row['customer']), [])
        if len(matches) != 1:
            raise HTTPException(409, f"판매처 연결을 확인해 주세요: {row['customer']}")
        if row['part_no'] not in items:
            registered = db.query(ItemMasterModel).filter(ItemMasterModel.part_no == row['part_no']).first()
            state = '품목 미등록' if not registered else f'자재유형 {registered.material_type}, 사용여부 {registered.is_active}'
            raise HTTPException(409, f"{row['part_no']}: {state}. 기초정보 품목정보에서 사용 중인 완제품/반제품인지 확인해 주세요. 자료는 변경되지 않았습니다.")
        row['customer_id'] = matches[0].id
        row['item_id'] = items[row['part_no']].id
    marker = 'UNSOLD-XLSX:' + hashlib.sha256(data).hexdigest()
    if db.query(SalesOrderMaster.id).filter(SalesOrderMaster.po_no == marker).first():
        raise HTTPException(409, '이미 등록한 파일입니다. 중복 등록할 수 없습니다.')
    orders = db.query(SalesOrderMaster).filter(SalesOrderMaster.status.in_(['ORDERED', 'PARTIAL'])).order_by(SalesOrderMaster.id).all()
    snapshot = [(o.id, o.status, [(i.id, i.order_qty, i.shipped_qty) for i in o.items]) for o in orders]
    fingerprint = hashlib.sha256(json.dumps([marker, rows, snapshot], sort_keys=True).encode()).hexdigest()
    return rows, items, orders, marker, fingerprint


async def upload_data(file):
    data = await file.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(413, '파일은 5MB 이하로 선택해 주세요.')
    return data


@router.post('/api/sales/unsold/import/preview')
async def preview(file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(require_admin_user)):
    rows, items, orders, marker, fingerprint = prepare(db, await upload_data(file))
    return dict(rows=rows, count=len(rows), qty=sum(r['qty'] for r in rows),
                close_count=len(orders), close_qty=sum(max(i.order_qty-i.shipped_qty, 0) for o in orders for i in o.items), fingerprint=fingerprint)


@router.post('/api/sales/unsold/import/apply')
async def apply(file: UploadFile = File(...), fingerprint: str = Form(...), db: Session = Depends(get_db), user=Depends(require_admin_user)):
    try:
        data = await upload_data(file)
        if db.get_bind().dialect.name == 'sqlite':
            db.execute(text('BEGIN IMMEDIATE'))
        rows, items, orders, marker, current = prepare(db, data)
        if fingerprint != current:
            raise HTTPException(409, '자료나 수주 상태가 변경되었습니다. 다시 미리보기를 실행해 주세요.')
        for order in orders:
            order.status = 'CANCELLED'
            order.note = (order.note or '') + '\n미판매 잔량 마감: 엑셀 기준 재등록. 원수량·출고 이력 보존.'
        for row in rows:
            master = items[row['part_no']]
            order = SalesOrderMaster(order_no=_next_no(db, SalesOrderMaster, SalesOrderMaster.order_no, 'SO', row['order_date']),
                order_date=row['order_date'], customer_id=row['customer_id'],
                customer_name=row['customer'], status='ORDERED', order_type='NORMAL', transaction_type='PAID',
                po_no=marker, note='이카운트 미판매 잔량 기준 등록 / 원본 ' + row['source'], created_by=_username(user))
            order.items.append(SalesOrderItem(item_id=master.id, part_no=master.part_no, part_name=master.part_name,
                order_qty=row['qty'], shipped_qty=0, unit=master.unit or 'EA', delivery_date=row['due'], note=row['note'], status='WAITING'))
            db.add(order)
            db.flush()
        db.commit()
        return dict(count=len(rows), qty=sum(r['qty'] for r in rows), closed=len(orders))
    except Exception:
        db.rollback()
        raise
