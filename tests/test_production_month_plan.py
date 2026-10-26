from datetime import datetime,date
from test_production_sync import setup
from test_production_run_registration import prepare
from routers.production_month_plan import router
from routers import production_month_plan as api
from routers.equipment_oee import router as oee_router
from models.production import ProductionWorkOrder,ProductionPerformance
from models.production_cycle_time import ProductionCycleTime
from models.production_month_plan import ProductionMonthPlanDay,ProductionEquipmentSchedule
from services.equipment_oee import calculate


def initialize(setup,monkeypatch):
    order,_,worker,equipment,other=prepare(setup)
    setup.app.include_router(router);setup.app.include_router(oee_router)
    class FixedDatetime(datetime):
        @classmethod
        def now(cls,tz=None):return cls(2026,10,12,12)
    monkeypatch.setattr(api,'datetime',FixedDatetime)
    with setup.sessions() as db:item=db.get(ProductionWorkOrder,order).item_id
    return order,item,equipment,other


def value(item,equipment,ct=30):
    return dict(month='2026-10',process_code='LT',lines=[dict(item_id=item,equipment_id=equipment,cycle_seconds=ct,expected_oee=.95,quantities={f'2026-10-{i:02d}':100 if i in (6,12) else 0 for i in range(1,32)})])


def test_excel_matches_saved_status_and_rejects_invalid_range(setup,monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    _,item,equipment,_=initialize(setup,monkeypatch)
    assert setup.client.put('/api/production/month-plan',json=value(item,equipment)).status_code==200
    params={'start_date':'2026-10-01','end_date':'2026-10-31','process_code':'LT'}
    response=setup.client.get('/api/production/month-plan/excel',params=params)
    assert response.status_code==200,response.text
    book=load_workbook(BytesIO(response.content))
    assert book.sheetnames==['공정별 요약','일자별 생산계획']
    sheet=book.worksheets[1]
    assert sheet['K3'].value==200 and sheet['Q3'].value==100 and sheet['W3'].value==100
    assert sheet['F3'].value==.95 and sheet['F3'].number_format=='0.00%'
    assert sheet['K4'].value==200 and sheet.freeze_panes=='L3'
    assert book.worksheets[0].cell(4,5).value==200
    params['start_date']='2026-10-12';params['end_date']='2026-10-12'
    day=load_workbook(BytesIO(setup.client.get('/api/production/month-plan/excel',params=params).content)).worksheets[1]
    assert day['K3'].value==100 and day['L3'].value==100
    params['end_date']='2026-10-01'
    assert setup.client.get('/api/production/month-plan/excel',params=params).status_code==422


def test_plan_ct_updates_whole_month_and_can_be_restored_without_changing_master(setup,monkeypatch):
    order,item,equipment,_=initialize(setup,monkeypatch)
    response=setup.client.put('/api/production/month-plan',json=value(item,equipment));assert response.status_code==200,response.text
    response=setup.client.put('/api/production/month-plan',json=value(item,equipment,40));assert response.status_code==200,response.text
    with setup.sessions() as db:
        rows={d.work_date:d for d in db.query(ProductionMonthPlanDay)}
        assert rows['2026-10-06'].cycle_seconds==40
        assert rows['2026-10-12'].cycle_seconds==40
        db.add(ProductionCycleTime(item_id=item,seconds=80))
        db.add(ProductionPerformance(work_order_id=order,performance_date='2026-10-06',process_code='LT',shift_type='DAY',equipment_id=equipment,good_qty=100,defect_qty=0,setup_qty=0));db.commit()
        before=calculate(db,date(2026,10,6),date(2026,10,6),now=datetime(2026,10,20))
        db.get(ProductionCycleTime,item).seconds=99;db.commit()
        after=calculate(db,date(2026,10,6),date(2026,10,6),now=datetime(2026,10,20))
        target=lambda r:next(g for g in r['items'] if g['key']==str(equipment))
        assert target(before)['oee'] is not None
        assert target(before)['oee']==target(after)['oee']
        assert target(after)['parts'][0]['cycle_seconds']==40
    # Restoring the original C/T rewrites all dates, with no effective-date requirement.
    assert setup.client.put('/api/production/month-plan',json=value(item,equipment,30)).status_code==200
    with setup.sessions() as db:
        assert {d.cycle_seconds for d in db.query(ProductionMonthPlanDay)}=={30}
    report=setup.client.get('/api/production/month-plan/status',params={'start_date':'2026-10-01','end_date':'2026-10-31'}).json()
    assert report['total']==200 and len(report['rows'])==1
    assert report['rows'][0]['cycle_seconds']==30


def test_detail_quantities_determine_each_equipment_workdays(setup,monkeypatch):
    _,item,equipment,other=initialize(setup,monkeypatch)
    assert setup.client.put('/api/production/month-plan',json=value(item,equipment)).status_code==200
    with setup.sessions() as db:
        result=calculate(db,date(2026,10,12),date(2026,10,12),now=datetime(2026,10,20))
        assert next(r for r in result['items'] if r['key']==str(equipment))['planned_seconds']==24*3600
        assert next(r for r in result['items'] if r['key']==str(other))['planned_seconds']==0
        assert db.query(ProductionEquipmentSchedule).count()==0
    empty=value(item,equipment)
    empty['lines'][0]['quantities']['2026-10-12']=0
    assert setup.client.put('/api/production/month-plan',json=empty).status_code==200
    with setup.sessions() as db:
        result=calculate(db,date(2026,10,12),date(2026,10,12),now=datetime(2026,10,20))
        assert all(r['planned_seconds']==0 for r in result['items'])



def test_duplicate_plan_rows_and_invalid_quantities_are_atomic(setup,monkeypatch):
    _,item,equipment,_=initialize(setup,monkeypatch)
    payload=value(item,equipment);payload['lines']*=2
    assert setup.client.put('/api/production/month-plan',json=payload).status_code==422
    payload=value(item,equipment);payload['lines'][0]['quantities']['2026-10-01']=-1
    assert setup.client.put('/api/production/month-plan',json=payload).status_code==422
    with setup.sessions() as db:assert db.query(ProductionMonthPlanDay).count()==0
    assert setup.client.get('/api/production/equipment/oee',params={'start_date':'2026-10-01','end_date':'2026-10-31','process_code':'DOT'}).status_code==422


def test_plan_references_are_reported_in_item_usage_and_block_equipment_deletion(setup,monkeypatch):
    from services.item_identity_service import item_usage_summary
    from routers import basic_info_equipment
    from test_production_sync import ADMIN
    _,item,equipment,_=initialize(setup,monkeypatch)
    assert setup.client.put('/api/production/month-plan',json=value(item,equipment)).status_code==200
    with setup.sessions() as db:
        assert any(row['table']=='production_plan_revision_references' for row in item_usage_summary(db,item))
    setup.app.include_router(basic_info_equipment.router)
    monkeypatch.setattr(basic_info_equipment,'get_current_user_optional',lambda request,db:ADMIN)
    assert setup.client.delete('/api/basic-info/equipment/'+str(equipment)).status_code==409


def test_multiple_detail_rows_share_one_equipment_day_and_ignore_old_calendar(setup,monkeypatch):
    from test_production_sync import add_item
    from models.oee_schedule import OeeScheduleDay
    _,item,equipment,other=initialize(setup,monkeypatch)
    with setup.sessions() as db:
        second=add_item(db,part='SECOND-A',code='LT')
        db.add(ProductionEquipmentSchedule(equipment_id=other,work_date='2026-10-12',working=True,start_time='08:00',end_time='08:00',next_day=True))
        db.add(OeeScheduleDay(process_code='LT',work_date='2026-10-13',working=False,start_time='08:00',end_time='08:00',next_day=True))
        db.commit()
    plan=value(item,equipment)
    quantities={f'2026-10-{i:02d}':0 for i in range(1,32)}
    plan['lines'][0]['quantities']=quantities|{'2026-10-12':100}
    plan['lines'].append(dict(item_id=second,equipment_id=equipment,cycle_seconds=45,expected_oee=.95,quantities=quantities|{'2026-10-13':50}))
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    with setup.sessions() as db:
        result=calculate(db,date(2026,10,12),date(2026,10,13),now=datetime(2026,10,20))
        assert next(r for r in result['items'] if r['key']==str(equipment))['planned_seconds']==48*3600
        assert next(r for r in result['items'] if r['key']==str(other))['planned_seconds']==0
    summary=setup.client.get('/api/production/month-plan/status',params={'start_date':'2026-10-12','end_date':'2026-10-13'}).json()
    assert summary['total']==150 and len(summary['rows'])==2
    plan['lines'][0]['quantities']['2026-10-12']=0
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    with setup.sessions() as db:
        result=calculate(db,date(2026,10,12),date(2026,10,13),now=datetime(2026,10,20))
        assert next(r for r in result['items'] if r['key']==str(equipment))['planned_seconds']==24*3600


def test_work_time_bulk_save_shared_hours_snapshot_and_atomic_validation(setup,monkeypatch):
    from models.production_month_plan import ProductionPlanWorkTime
    from services import oee_schedule
    _,item,equipment,other=initialize(setup,monkeypatch)
    plan=value(item,equipment)
    plan['lines'][0]['quantities']['2026-10-16']=100
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    with setup.sessions() as db:
        assert db.query(ProductionPlanWorkTime).count()==31
        rows=oee_schedule.days(db,'LT',date(2026,10,16),date(2026,10,16),equipment)
        assert rows[0]['end_time']=='08:00'  # Friday also defaults to 24 hours.
    overrides=[dict(equipment_id=equipment,work_date=d,start_time='08:00',end_time='05:00',next_day=True) for d in ['2026-10-12','2026-10-16']]
    plan['work_times']=overrides
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    report=setup.client.get('/api/production/month-plan/status',params={'start_date':'2026-10-16','end_date':'2026-10-16'}).json()
    assert report['rows'][0]['upm_full']==report['rows'][0]['uph']*21
    with setup.sessions() as db:
        result=calculate(db,date(2026,10,16),date(2026,10,16),now=datetime(2026,10,20))
        assert next(r for r in result['items'] if r['key']==str(equipment))['planned_seconds']==21*3600
        assert next(r for r in result['items'] if r['key']==str(other))['planned_seconds']==0
    # A later default change or a quantity-only save cannot change saved shifts.
    original=oee_schedule.default_day
    monkeypatch.setattr(oee_schedule,'default_day',lambda d,h:original(d,h)|{'end_time':'04:00'})
    plan.pop('work_times')
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    loaded=setup.client.get('/api/production/month-plan',params={'month':'2026-10','process_code':'LT'}).json()
    rows={r['work_date']:r for r in loaded['schedules'][str(equipment)]}
    assert rows['2026-10-16']['end_time']=='05:00' and rows['2026-10-13']['end_time']=='08:00'
    # Invalid second override rolls back the valid first and quantity changes too.
    plan['lines'][0]['quantities']['2026-10-12']=999
    plan['work_times']=[overrides[0]|{'end_time':'06:00'},overrides[1]|{'next_day':False,'end_time':'07:00'}]
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==422
    with setup.sessions() as db:
        assert db.query(ProductionPlanWorkTime).filter_by(equipment_id=equipment,work_date='2026-10-12').one().end_time=='05:00'
        assert db.query(ProductionMonthPlanDay).filter_by(work_date='2026-10-12').one().quantity==100
    for bad in [overrides[0]|{'work_date':'2026-11-01'},overrides[0]|{'start_time':'24:00'},overrides[0]|{'equipment_id':other}]:
        plan['work_times']=[bad]
        assert setup.client.put('/api/production/month-plan',json=plan).status_code==422


def test_existing_plan_time_migration_preserves_past_friday_and_is_idempotent(setup,monkeypatch):
    from models.production_month_plan import ProductionPlanWorkTime
    from models.production_plan_time_migration import ensure_production_plan_work_times
    _,item,equipment,_=initialize(setup,monkeypatch)
    assert setup.client.put('/api/production/month-plan',json=value(item,equipment)).status_code==200
    with setup.sessions() as db:
        db.query(ProductionPlanWorkTime).delete();db.commit()
        engine=db.get_bind()
    ensure_production_plan_work_times(engine,today=date(2026,10,12))
    with setup.sessions() as db:
        past=db.query(ProductionPlanWorkTime).filter_by(work_date='2026-10-09').one()
        assert past.end_time=='05:00'
        assert db.query(ProductionPlanWorkTime).filter_by(work_date='2026-10-16').one().end_time=='08:00'
        past.end_time='06:00';db.commit()
    ensure_production_plan_work_times(engine,today=date(2026,11,1))
    with setup.sessions() as db:
        assert db.query(ProductionPlanWorkTime).count()==31
        assert db.query(ProductionPlanWorkTime).filter_by(work_date='2026-10-09').one().end_time=='06:00'
        assert db.query(ProductionPlanWorkTime).filter_by(work_date='2026-10-16').one().end_time=='08:00'


def test_explicit_row_delete_removes_plan_days_but_keeps_performance_and_other_rows(setup,monkeypatch):
    from models.production_month_plan import ProductionMonthPlanLine,ProductionPlanWorkTime
    order,item,equipment,other=initialize(setup,monkeypatch)
    plan=value(item,equipment)
    plan['lines'].append(plan['lines'][0]|{'equipment_id':other})
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    with setup.sessions() as db:
        doomed=db.query(ProductionMonthPlanLine).filter_by(equipment_id=equipment).one().id
        remaining=db.query(ProductionMonthPlanLine).filter_by(equipment_id=other).one().id
        db.add(ProductionPerformance(work_order_id=order,performance_date='2026-10-06',process_code='LT',shift_type='DAY',equipment_id=equipment,good_qty=100,defect_qty=0,setup_qty=0));db.commit()
    # Omitted rows alone are never deleted (older clients retain their behavior).
    partial=plan|{'lines':[plan['lines'][1]]}
    assert setup.client.put('/api/production/month-plan',json=partial).status_code==200
    with setup.sessions() as db:assert db.query(ProductionMonthPlanLine).count()==2
    deletion=partial|{'deleted_line_ids':[doomed]}
    assert setup.client.put('/api/production/month-plan',json=deletion).status_code==200
    with setup.sessions() as db:
        assert db.get(ProductionMonthPlanLine,doomed) is None
        assert db.get(ProductionMonthPlanLine,remaining) is not None
        assert db.query(ProductionMonthPlanDay).count()==31
        assert db.query(ProductionPerformance).count()==1
        assert db.query(ProductionPlanWorkTime).count()==62
    loaded=setup.client.get('/api/production/month-plan',params={'month':'2026-10','process_code':'LT'}).json()
    assert [(r['equipment_id'],r['item_id']) for r in loaded['lines']]==[(other,item)]
    report=setup.client.get('/api/production/month-plan/status',params={'start_date':'2026-10-01','end_date':'2026-10-31'}).json()
    assert report['total']==200 and len(report['rows'])==1
    assert setup.client.put('/api/production/month-plan',json=partial|{'lines':[],'deleted_line_ids':[remaining]}).status_code==200
    with setup.sessions() as db:assert db.query(ProductionMonthPlanDay).count()==0


def test_delete_scope_validation_and_invalid_save_are_atomic(setup,monkeypatch):
    from models.production_month_plan import ProductionMonthPlanLine
    _,item,equipment,_=initialize(setup,monkeypatch)
    plan=value(item,equipment)
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    with setup.sessions() as db:line_id=db.query(ProductionMonthPlanLine).one().id
    for bad in [plan|{'month':'2026-11','lines':[],'deleted_line_ids':[line_id]},plan|{'process_code':'TP','lines':[],'deleted_line_ids':[line_id]},plan|{'lines':[],'deleted_line_ids':[line_id+100]}]:
        assert setup.client.put('/api/production/month-plan',json=bad).status_code==422
    bad=plan|{'lines':[],'deleted_line_ids':[line_id],'work_times':[dict(equipment_id=equipment,work_date='2026-10-12',start_time='08:00',end_time='05:00',next_day=False)]}
    assert setup.client.put('/api/production/month-plan',json=bad).status_code==422
    with setup.sessions() as db:
        assert db.query(ProductionMonthPlanLine).count()==1 and db.query(ProductionMonthPlanDay).count()==31
    from types import SimpleNamespace
    from core.security import get_current_user
    import json
    setup.app.dependency_overrides[get_current_user]=lambda:SimpleNamespace(username='reader',role='USER',permissions=json.dumps({'menu_access':{'/production/plans':'READ'}}))
    assert setup.client.put('/api/production/month-plan',json=plan|{'lines':[],'deleted_line_ids':[line_id]}).status_code==403


def test_tapping_planned_exclusion_updates_capacity_and_oee_without_changing_shift_bounds(setup,monkeypatch):
    from test_production_sync import add_item
    from models.equipment import EquipmentMaster
    from models.production_month_plan import ProductionPlanWorkTime
    import pytest
    initialize(setup,monkeypatch)
    with setup.sessions() as db:
        item=add_item(db,part='TAPPING-B',code='TP')
        equipment=EquipmentMaster(equipment_code='TP-01',equipment_name='탭핑 1호기',process_code='TP',machine_no='1',is_active='Y')
        db.add(equipment);db.flush()
        order=ProductionWorkOrder(work_order_no='TAP-PLAN',order_date='2026-10-12',item_id=item,part_no='TAPPING-B',order_qty=1000)
        db.add(order);db.flush()
        db.add(ProductionPerformance(work_order_id=order.id,performance_date='2026-10-12',process_code='TP',equipment_id=equipment.id,good_qty=900,defect_qty=100,setup_qty=0));db.commit()
        equip_id=equipment.id
    plan=value(item,equip_id)|{'process_code':'TP'}
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    loaded=setup.client.get('/api/production/month-plan',params={'month':'2026-10','process_code':'TP'}).json()
    day=next(d for d in loaded['schedules'][str(equip_id)] if d['work_date']=='2026-10-12')
    assert day['start_time']==day['end_time']=='08:00' and day['next_day'] and day['excluded_minutes']==210
    report=setup.client.get('/api/production/month-plan/status',params={'start_date':'2026-10-12','end_date':'2026-10-12','process_code':'TP'}).json()['rows'][0]
    assert report['upm_full']==2337 and report['upm_day']==1178 and report['upd']==1178
    with setup.sessions() as db:
        result=calculate(db,date(2026,10,12),date(2026,10,12),now=datetime(2026,10,14),process_code='TP')
        assert result['planned_seconds']==pytest.approx(73800)
        assert result['items'][0]['oee']==pytest.approx(900*30/73800)
        partial=calculate(db,date(2026,10,12),date(2026,10,12),now=datetime(2026,10,12,12),process_code='TP')
        assert partial['planned_seconds']==pytest.approx(12300)
    plan['work_times']=[dict(equipment_id=equip_id,work_date='2026-10-12',start_time='08:00',end_time='08:00',next_day=True,excluded_minutes=60)]
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    with setup.sessions() as db:
        assert db.query(ProductionPlanWorkTime).filter_by(equipment_id=equip_id,work_date='2026-10-12').one().excluded_minutes==60
        assert calculate(db,date(2026,10,12),date(2026,10,12),now=datetime(2026,10,14),process_code='TP')['planned_seconds']==23*3600
    for bad in [-1,1440]:
        plan['work_times'][0]['excluded_minutes']=bad
        assert setup.client.put('/api/production/month-plan',json=plan).status_code==422
    plan['work_times'][0]|={'excluded_minutes':190,'next_day':False,'end_time':'10:00'}
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==422
    with setup.sessions() as db:assert db.query(ProductionPlanWorkTime).filter_by(equipment_id=equip_id,work_date='2026-10-12').one().excluded_minutes==60
    plan['work_times'][0]|={'excluded_minutes':90,'next_day':True,'end_time':'08:00'}
    assert setup.client.put('/api/production/month-plan',json=plan).status_code==200
    two_days=setup.client.get('/api/production/month-plan/status',params={'start_date':'2026-10-01','end_date':'2026-10-31','process_code':'TP'}).json()['rows'][0]
    assert two_days['upm_full']==4902 and two_days['upm_day']==2460


def test_legacy_tapping_exclusion_migration_is_once_only_and_keeps_shorter_custom_shifts(setup,monkeypatch):
    from sqlalchemy import text
    from models.equipment import EquipmentMaster
    from models.production_month_plan import ProductionPlanWorkTime
    from models.production_plan_time_migration import ensure_production_plan_work_times
    _,_,lathe,_=initialize(setup,monkeypatch)
    with setup.sessions() as db:
        tap=EquipmentMaster(equipment_code='TP-01',equipment_name='탭핑 1호기',process_code='TP',machine_no='1',is_active='Y');db.add(tap);db.flush()
        db.add_all([ProductionPlanWorkTime(equipment_id=equip,work_date=day,start_time='08:00',end_time=end,next_day=True) for equip,day,end in [(tap.id,'2026-10-12','08:00'),(tap.id,'2026-10-13','05:00'),(lathe,'2026-10-12','08:00')]])
        db.commit();tap_id=tap.id
    with setup.engine.begin() as connection:connection.execute(text('ALTER TABLE production_plan_work_times DROP COLUMN excluded_minutes'))
    ensure_production_plan_work_times(setup.engine,today=date(2026,10,13))
    with setup.sessions() as db:
        full=db.query(ProductionPlanWorkTime).filter_by(equipment_id=tap_id,work_date='2026-10-12').one()
        assert full.excluded_minutes==210 and full.end_time=='08:00'
        assert db.query(ProductionPlanWorkTime).filter_by(equipment_id=tap_id,work_date='2026-10-13').one().excluded_minutes==0
        assert db.query(ProductionPlanWorkTime).filter_by(equipment_id=lathe).one().excluded_minutes==0
        full.excluded_minutes=0;db.commit()
    ensure_production_plan_work_times(setup.engine,today=date(2026,10,13))
    with setup.sessions() as db:assert db.query(ProductionPlanWorkTime).filter_by(equipment_id=tap_id,work_date='2026-10-12').one().excluded_minutes==0


def test_old_190_minute_defaults_correct_once_and_keep_custom_hours(setup,monkeypatch):
    from models.equipment import EquipmentMaster
    from models.production_month_plan import ProductionPlanWorkTime
    from models.production_plan_time_migration import ensure_production_plan_work_times
    _,_,lathe,_=initialize(setup,monkeypatch)
    with setup.sessions() as db:
        tap=EquipmentMaster(equipment_code='TP-OLD',equipment_name='탭핑 1호기',process_code='TP',machine_no='1');db.add(tap);db.flush()
        for equip,day,end,excluded in [(tap.id,'2026-10-01','08:00',190),(tap.id,'2026-10-02','08:00',60),(tap.id,'2026-10-03','05:00',190),(lathe,'2026-10-02','05:00',0)]:
            db.add(ProductionPlanWorkTime(equipment_id=equip,work_date=day,start_time='08:00',end_time=end,next_day=True,excluded_minutes=excluded))
        db.commit();tap_id=tap.id
    ensure_production_plan_work_times(setup.engine,today=date(2026,10,13))
    with setup.sessions() as db:
        times={r.work_date:r for r in db.query(ProductionPlanWorkTime).filter_by(equipment_id=tap_id)}
        assert times['2026-10-01'].excluded_minutes==210
        assert times['2026-10-02'].excluded_minutes==60
        assert times['2026-10-03'].excluded_minutes==190
        assert db.query(ProductionPlanWorkTime).filter_by(equipment_id=lathe,work_date='2026-10-02').one().end_time=='05:00'
        times['2026-10-01'].excluded_minutes=190;db.commit()
    ensure_production_plan_work_times(setup.engine,today=date(2026,10,13))
    with setup.sessions() as db:
        assert db.query(ProductionPlanWorkTime).filter_by(equipment_id=tap_id,work_date='2026-10-01').one().excluded_minutes==190


def test_other_registered_process_items_are_not_searchable_or_addable_to_machining_plans(setup,monkeypatch):
    from test_production_sync import add_item
    from models.equipment import EquipmentMaster
    from models.models import ItemBomModel
    from models.production_month_plan import ProductionMonthPlanLine
    initialize(setup,monkeypatch)
    with setup.sessions() as db:
        machine=EquipmentMaster(equipment_code='TP-AG-TEST',equipment_name='탭핑 검증',process_code='TP',machine_no='1',is_active='Y')
        db.add(machine);db.flush();machine_id=machine.id
        valid=add_item(db,part='310068-B',code='TP')
        child=add_item(db,part='AG-TEST-RAW')
        plated=[]
        for part in ['310068-Ag','310069-AG','310070-ag']:
            item=add_item(db,part=part,code='PL');plated.append(item)
            # A tapping BOM must not override the item master plating registration.
            db.add(ItemBomModel(parent_item_id=item,child_item_id=child,parent_part_no=part,child_part_no='AG-TEST-RAW',process_code='TP',quantity=1,unit='EA',created_at='2026-01-01'))
        db.commit()
    candidates=setup.client.get('/api/production/month-plan/options').json()['items']
    assert any(r['id']==valid and 'TP' in r['process_codes'] for r in candidates)
    assert not {r['id'] for r in candidates}&set(plated)
    for item in plated:
        body=value(item,machine_id)|{'process_code':'TP'}
        assert setup.client.put('/api/production/month-plan',json=body).status_code==422
        generation=setup.client.get('/api/production/month-plan/revisions',params={'month':'2026-10'}).json()['generation']
        assert setup.client.put('/api/production/month-plan/revisions/0',json=body|{'generation':generation}).status_code==422
    with setup.sessions() as db:assert db.query(ProductionMonthPlanLine).count()==0
