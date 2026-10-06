"""Company-sheet imports must validate completely before touching an FMEA draft."""
import io
import os
import tempfile
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from services.fmea_excel_import import FmeaImportError, MAX_FILE_BYTES, parse_company_fmea


def company_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "갑지"
    sheet['A1'] = '기본정보는 불러오지 않음'
    sheet['B6'] = '다른 품번이어도 상단 정보는 가져오지 않음'
    labels = {1:'공정의 기능',2:'잠재적 고장형태',3:'고장의 잠재적 영향',4:'심각도',5:'특별특성',
              6:'고장의 잠재적 원인',8:'발생도',9:'현공정 관리',11:'검출도',12:'R. P. N.',
              13:'권고 조치사항',15:'완료예정일',16:'조치결과'}
    for col, value in labels.items(): sheet.cell(8, col, value)
    for col, value in {16:'조치내용',17:'심각도',18:'발생도',19:'검출도',20:'R. P. N.'}.items(): sheet.cell(9, col, value)
    sheet['I10']='예방';sheet['J10']='검출'
    sheet.merge_cells('A11:A13');sheet['A11']='10\n입고검사'
    sheet['A14']='20 가공'
    for first, last, mode in [(11,12,'치수불량'),(13,13,'외관불량'),(14,14,'BURR')]:
        for col in (2,3,4,8,11,12,15,16,17,18,19,20):
            if first != last: sheet.merge_cells(start_row=first,start_column=col,end_row=last,end_column=col)
        sheet.cell(first,2,mode);sheet.cell(first,3,'사용불가')
        for col, value in [(4,3),(8,2),(11,4)]: sheet.cell(first,col,value)
        sheet.cell(first,12,f'=K{first}*H{first}*D{first}')
        sheet.merge_cells(start_row=first,start_column=13,end_row=last,end_column=14)
        sheet.cell(first,13,'N/A')
        for col in (15,16,17,18,19,20): sheet.cell(first,col,'N/A')
        for row in range(first,last+1):
            sheet.merge_cells(start_row=row,start_column=6,end_row=row,end_column=7)
            sheet.cell(row,6,'원인 '+str(row));sheet.cell(row,9,'예방 '+str(row));sheet.cell(row,10,'검출 '+str(row))
    sheet['A15']='RPN (Risk Priority Number) 위험순위도'
    workbook.create_sheet('을지')['A1']='사용하지 않는 잘못된 항목과 공정'
    workbook.create_sheet('병지')['A1']='사용하지 않는 잘못된 항목과 공정'
    return workbook


def content(workbook):
    target=io.BytesIO();workbook.save(target);return target.getvalue()


def steps():
    return [SimpleNamespace(id=1,step_no='10',step_name='입고검사'),SimpleNamespace(id=2,step_no='20',step_name='가공')]


def test_preserves_merged_multiline_analysis_and_ignores_headers_and_other_sheets():
    rows=parse_company_fmea(content(company_workbook()),steps())
    assert len(rows)==3
    assert [r['flow_step_id'] for r in rows]==[1,1,2]
    assert rows[0]['causes']=='원인 11\n원인 12'
    assert rows[0]['effects']=='사용불가'
    assert rows[0]['prevention_controls']=='예방 11\n예방 12'
    assert rows[0]['detection_controls']=='검출 11\n검출 12'
    assert rows[0]['action_not_applicable'] is True
    assert rows[0]['function_text']=='입고검사'
    assert not any('company' in r or 'document_no' in r for r in rows)


@pytest.mark.parametrize('cell,value,message',[
    ('A14','20 다른공정','공정명 불일치'),('A14','30 가공','공정번호'),
    ('B8','잘못된 항목','항목명'),('D11',11,'1~10'),('D11','9'*5000,'1~10'),('L11',999,'RPN 불일치'),
    ('F11','=1+1','수식'),('L11','=K14*H14*D14','같은 분석행'),
    ('M13','SPC','혼재'),('U13','누락되면 안 되는 내용','밖에 데이터'),
])
def test_invalid_cells_reject_entire_file(cell,value,message):
    workbook=company_workbook();workbook['갑지'][cell]=value
    with pytest.raises(FmeaImportError) as result:
        parse_company_fmea(content(workbook),steps())
    assert any(message in error['message'] for error in result.value.errors)


def test_order_and_missing_stage_rejected():
    workbook=company_workbook()
    with pytest.raises(FmeaImportError) as result:
        parse_company_fmea(content(workbook),list(reversed(steps())))
    assert any('순서' in x['message'] for x in result.value.errors)
    with pytest.raises(FmeaImportError): parse_company_fmea(content(workbook),steps()+[SimpleNamespace(id=3,step_no='30',step_name='포장')])


def test_partial_date_rejected_and_valid_action_scores_accepted():
    workbook=company_workbook();sheet=workbook['갑지']
    for cell,value in [('M13','SPC 관리'),('O13',21.11),('P13','관리도'),('Q13',3),('R13',2),('S13',2),('T13','=Q13*R13*S13')]: sheet[cell]=value
    with pytest.raises(FmeaImportError) as result:
        parse_company_fmea(content(workbook),steps())
    assert any(x['cell']=='갑지!O13' and '연·월·일' in x['message'] for x in result.value.errors)
    sheet['O13']=date(2021,11,30)
    rows=parse_company_fmea(content(workbook),steps())
    assert rows[1]['target_date']=='2021-11-30' and rows[1]['new_detection']==2
    assert rows[1]['action_not_applicable'] is False


def test_missing_sheet_corrupt_file_and_size_rejected():
    workbook=company_workbook();workbook['갑지'].title='다른 시트'
    for data in (content(workbook),b'not an xlsx',b'x'*(MAX_FILE_BYTES+1)):
        with pytest.raises(FmeaImportError): parse_company_fmea(data,steps())


def test_cross_field_merge_and_orphan_data_rejected():
    workbook=company_workbook();workbook['갑지'].merge_cells('I11:J11')
    with pytest.raises(FmeaImportError): parse_company_fmea(content(workbook),steps())
    workbook=company_workbook();workbook['갑지'].unmerge_cells('B11:B12')
    with pytest.raises(FmeaImportError) as result: parse_company_fmea(content(workbook),steps())
    assert any('고장형태' in x['message'] for x in result.value.errors)


# Isolate model import side effects from the tracked and live development databases.
_bootstrap=tempfile.TemporaryDirectory()
os.environ['DATABASE_URL']='sqlite:///'+str(Path(_bootstrap.name)/'bootstrap.db')
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select, func
from sqlalchemy.orm import sessionmaker
from core.database import Base, get_db, engine as bootstrap_engine
from core.security import get_current_user
from models.models import CommonCodeModel, ItemMasterModel, UserModel
from models.fmea import FmeaDocument, FmeaRevision, FmeaRow
from models.process_flow import ProcessFlowRevision, ProcessFlowStep, ProcessFlowStepKey
from routers.process_fmea import router


@pytest.fixture(scope='session',autouse=True)
def cleanup_bootstrap():
    yield
    bootstrap_engine.dispose();_bootstrap.cleanup()


@pytest.fixture
def api(tmp_path):
    engine=create_engine('sqlite:///'+str(tmp_path/'test.db'),connect_args={'check_same_thread':False})
    Base.metadata.create_all(engine)
    factory=sessionmaker(bind=engine,expire_on_commit=False)
    with factory() as db:
        admin=UserModel(id=1,username='tester',password_hash='not-a-login',name='테스터',role='ADMIN',created_at='2026-10-06')
        db.add(admin)
        db.add(CommonCodeModel(group_code='MATERIAL_TYPE',code='FINISHED',group_name='재질',code_name='완제품',created_at='2026-10-06'))
        db.add(ItemMasterModel(id=1,part_no='310061',part_name='완제품',account_type='PROD',material_type='FINISHED',created_at='2026-10-06'))
        db.add(ProcessFlowRevision(id=1,item_id=1,revision_code='REV.0',sequence=1,version=1,status='CURRENT',part_no_snapshot='310061',part_name_snapshot='완제품',created_by_id=1,created_by='테스터'))
        db.flush()
        for i,step in enumerate(steps(),1):
            db.add(ProcessFlowStepKey(id=i,item_id=1));db.flush()
            db.add(ProcessFlowStep(id=i,revision_id=1,step_key_id=i,sort_order=i,step_no=step.step_no,step_name=step.step_name))
        db.add(FmeaDocument(id=1,item_id=1,document_no='TEST',created_by_id=1,created_by='테스터'));db.flush()
        db.add(FmeaRevision(id=1,document_id=1,revision_code='REV.0',sequence=1,status='DRAFT',version=1,flow_revision_id=1,part_no_snapshot='310061',part_name_snapshot='완제품',prepared_by='작성자 유지',date_prepared=date(2026,10,6),company='회사 유지',created_by_id=1,created_by='테스터'));db.flush()
        db.add(FmeaRow(id=1,revision_id=1,sort_order=1,flow_step_id=1,function_text='기존 기능',failure_mode='기존 분석행',effects='영향',causes='원인'))
        db.commit()
    app=FastAPI();app.include_router(router)
    def sessions():
        with factory() as db: yield db
    app.dependency_overrides[get_db]=sessions
    app.dependency_overrides[get_current_user]=lambda:admin
    with TestClient(app) as client: yield client,factory,admin
    engine.dispose()


def upload(client,workbook=None,**changes):
    fields={'item_id':'1','flow_revision_id':'1','flow_version':'1','revision_id':'1','revision_version':'1'}
    fields.update(changes)
    return client.post('/api/process-fmea/import-excel',data=fields,files={'file':('company.xlsx',content(workbook or company_workbook()),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})


def test_api_returns_validated_rows_without_writing_any_data(api):
    client,factory,_=api
    response=upload(client)
    assert response.status_code==200,response.text
    assert response.json()['row_count']==3
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(FmeaRow))==1
        assert db.get(FmeaRow,1).failure_mode=='기존 분석행'
        assert db.get(FmeaRevision,1).version==1
        assert db.get(FmeaRevision,1).company=='회사 유지'


def test_api_rejects_mismatch_without_partial_replacement(api):
    client,factory,_=api
    workbook=company_workbook();workbook['갑지']['D14']=99
    response=upload(client,workbook)
    assert response.status_code==422
    assert any(x['cell']=='갑지!D14' for x in response.json()['detail']['errors'])
    with factory() as db:
        assert db.get(FmeaRow,1).failure_mode=='기존 분석행'
        assert db.get(FmeaRevision,1).version==1
        assert db.scalar(select(func.count()).select_from(FmeaRow))==1


@pytest.mark.parametrize('case',['permission','flow_version','revision_version','current_revision','other_item','unauthenticated'])
def test_api_guards_permissions_identity_status_and_versions(api,case):
    client,factory,admin=api
    kwargs={};expected=409
    if case=='permission': admin.role='USER';admin.permissions='{}';expected=403
    elif case=='unauthenticated':
        from fastapi import HTTPException
        def missing(): raise HTTPException(401,'로그인이 필요합니다.')
        client.app.dependency_overrides[get_current_user]=missing;expected=401
    elif case=='flow_version': kwargs['flow_version']='2'
    elif case=='revision_version': kwargs['revision_version']='2'
    elif case=='other_item': kwargs['item_id']='999';expected=404
    else:
        with factory() as db: db.get(FmeaRevision,1).status='CURRENT';db.commit()
    assert upload(client,**kwargs).status_code==expected


def test_import_save_preserves_header_retires_old_rows_and_blocks_changed_flow(api):
    client,factory,_=api
    imported=upload(client).json()
    header={'company':'회사 유지','prepared_by':'작성자 유지','date_prepared':'2026-10-06',
            'flow_revision_id':1,'import_flow_version':imported['flow_version'],'version':1,'rows':imported['rows']}
    with factory() as db: db.get(ProcessFlowRevision,1).version=2;db.commit()
    response=client.put('/api/process-fmea/revisions/1',json=header)
    assert response.status_code==409
    with factory() as db:
        assert db.get(FmeaRow,1).retired_at is None
        assert db.get(FmeaRevision,1).version==1
        db.get(ProcessFlowRevision,1).version=1;db.commit()
    response=client.put('/api/process-fmea/revisions/1',json=header)
    assert response.status_code==200,response.text
    with factory() as db:
        revision=db.get(FmeaRevision,1)
        assert revision.company=='회사 유지' and revision.prepared_by=='작성자 유지'
        assert revision.version==2
        assert db.get(FmeaRow,1).retired_at is not None
        assert db.scalar(select(func.count()).select_from(FmeaRow).where(FmeaRow.retired_at.is_(None)))==3


def test_api_rejects_wrong_file_extension_and_cross_item_flow(api):
    client,factory,_=api
    response=client.post('/api/process-fmea/import-excel',data={'item_id':'1','flow_revision_id':'1','flow_version':'1'},files={'file':('company.xls',content(company_workbook()))})
    assert response.status_code==422
    with factory() as db:
        db.add(ItemMasterModel(id=2,part_no='OTHER',part_name='다른 품목',account_type='PROD',material_type='FINISHED',created_at='2026-10-06'))
        db.commit()
    response=client.post('/api/process-fmea/import-excel',data={'item_id':'2','flow_revision_id':'1','flow_version':'1'},files={'file':('company.xlsx',content(company_workbook()))})
    assert response.status_code==422 and '같은 품목' in response.json()['detail']


def test_product_rpn_and_duplicate_stage_numbers():
    workbook=company_workbook();workbook['갑지']['L11']='=PRODUCT($D$11,$H$11,$K$11)'
    assert len(parse_company_fmea(content(workbook),steps()))==3
    with pytest.raises(FmeaImportError) as result:
        parse_company_fmea(content(workbook),steps()+[SimpleNamespace(id=3,step_no='10',step_name='다른 공정')])
    assert any('중복' in error['message'] for error in result.value.errors)
