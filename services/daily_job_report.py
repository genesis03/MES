import html,math,copy
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font,Border,Side

TEMPLATE=Path(__file__).resolve().parents[1]/'assets'/'DailyJobReport.xlsx'

def put(sheet,cell,value):
    target=sheet[cell];target.value=value
    if isinstance(value,str):target.data_type='s'
    alignment=copy.copy(target.alignment)
    alignment.wrap_text=True
    target.alignment=alignment


def time_text(value):
    text=str(value or '')
    if len(text)==14 and text.isdigit():return text[8:10]+':'+text[10:12]
    if len(text)>=16 and text[10] in [' ','T']:return text[11:16]
    return text


def build_workbook(rows):
    workbook=load_workbook(TEMPLATE);template=workbook.active
    groups=defaultdict(list)
    for row in rows:
        extra=row['supplement']
        key=(row['performance_date'],row['process_name'],row['equipment_name'],extra['shift'] or row['shift_name'])
        groups[key].append(row)
    page_no=0
    for key,items in sorted(groups.items()):
        for offset in range(0,len(items),2):
            pair=items[offset:offset+2]
            inspection_pages=max(1,max(math.ceil(len(row['supplement']['measurements'])/8) for row in pair))
            for section in range(inspection_pages):
                sheet=workbook.copy_worksheet(template);page_no+=1;sheet.title=f'Job{page_no}'
                for row_cells in sheet:
                    for cell in row_cells:
                        if cell.fill.patternType=='solid' and cell.fill.fgColor.type=='rgb' and cell.fill.fgColor.rgb=='FFFFFF00':
                            put(sheet,cell.coordinate,None)
                put(sheet,'C1',key[0]);put(sheet,'C2',key[2]);put(sheet,'C3',key[3]);put(sheet,'O5','총생산')
                # Keep the measurement area width while giving each sample equal space.
                left_width=sum(sheet.column_dimensions[col].width for col in ['D','E','F','G'])/3
                for col in ['D','E']:sheet.column_dimensions[col].width=left_width/2
                for col in ['F','G']:sheet.column_dimensions[col].width=left_width
                right_width=sum(sheet.column_dimensions[col].width for col in ['L','M','N'])/3
                for col in ['L','M','N']:sheet.column_dimensions[col].width=right_width
                # Replace each merged measurement cell with three sample columns.
                measurement_styles={r:copy.copy(sheet[f'D{r}']._style) for r in range(14,25)}
                for merged in ['D14:G15','L14:N15']+[f'D{r}:G{r}' for r in range(16,25)]+[f'L{r}:N{r}' for r in range(16,25)]:
                    sheet.unmerge_cells(merged)
                sheet.merge_cells('D14:E15');sheet.merge_cells('F14:F15');sheet.merge_cells('G14:G15')
                for col in ['L','M','N']:sheet.merge_cells(f'{col}14:{col}15')
                for i,col in enumerate(['D','F','G']):put(sheet,col+'14',f'측정 {i+1}');sheet[col+'14'].font=Font(size=7,bold=True)
                for i,col in enumerate(['L','M','N']):put(sheet,col+'14',f'측정 {i+1}');sheet[col+'14'].font=Font(size=7,bold=True)
                for r in range(16,25):sheet.merge_cells(f'D{r}:E{r}')
                for r in range(14,25):
                    for col in ['D','F','G','L','M','N']:
                        cell=sheet[f'{col}{r}']
                        if cell.__class__.__name__!='MergedCell':
                            cell._style=copy.copy(measurement_styles[r])
                            cell.border=Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
                for col in ['D','F','G','L','M','N']:
                    sheet[col+'14'].font=Font(name='맑은 고딕',size=7,bold=True)
                # Clear all template placeholder values in input areas.
                for line in [7,8]:
                    for col in ['D','G','I','M','O','Q','S','U','V','W','Z','AH']:put(sheet,f'{col}{line}',None)
                for col in ['D','L']:
                    put(sheet,col+'12',None);put(sheet,col+'13',None)
                for r in range(16,25):
                    for col in ['D','F','G','L','M','N','H','J','O','R']:put(sheet,f'{col}{r}',None)
                for idx,row in enumerate(pair):
                    extra=row['supplement'];line=7+idx
                    values={'D':extra['operator_name'] or row['operator_name'],'G':row['part_no'],'I':row['part_name'],
                        'M':extra['time'] or ' ~ '.join(time_text(v) for v in row['work_time'].split(' ~ ')),
                        'O':row['total_qty'],'Q':row['setup_qty'],'W':row['defect_qty'],
                        'Z':row['material_lots'],'AH':row['output_lot_no']}
                    for col,value in values.items():put(sheet,f'{col}{line}',value)
                    left=idx==0;prefix='D' if left else 'L'
                    put(sheet,prefix+'12',row['part_no']);put(sheet,prefix+'13',extra['inspection_time'])
                    put(sheet,('C' if left else 'K')+'11',extra['inspection_kind'] or 'SET-UP 검사 / 자주검사')
                    for n,measurement in enumerate(extra['measurements'][section*8:section*8+8],16):
                        put(sheet,('C' if left else 'K')+str(n),measurement['label'] or str(section*8+n-15))
                        from openpyxl.comments import Comment
                        sheet[('C' if left else 'K')+str(n)].comment=Comment(measurement['spec'] or '기준 미입력','MES')
                        for col,value in zip(['D','F','G'] if left else ['L','M','N'],measurement['values']):put(sheet,col+str(n),value)
                        put(sheet,('H' if left else 'O')+str(n),measurement['worker_result']);put(sheet,('J' if left else 'R')+str(n),measurement['quality_result'])
                    put(sheet,prefix+'24',extra['appearance'])
                for field,cell in [('cut_bars','AC12'),('raw_setup','AF12')]:
                    vals=[row['supplement'][field] for row in pair]
                    put(sheet,cell,sum(vals) if all(v is not None for v in vals) else None)
                ins=[r['supplement']['raw_input'] for r in pair];rem=[r['supplement']['raw_remaining'] for r in pair]
                put(sheet,'Y12',f'{sum(ins)} / {sum(rem)}' if all(v is not None for v in ins+rem) else None)
                defects=[(row,defect) for row in pair for defect in row.get('defects',[])]
                for r,(source,defect) in enumerate(defects[:4],15):
                    put(sheet,'Y'+str(r),source['part_no']);put(sheet,'AA'+str(r),defect['qty']);put(sheet,'AD'+str(r),defect['name'])
                if section>0:put(sheet,'F1',f'작업일보 — 검사 계속 {section+1}/{inspection_pages}')
                notes=[r['supplement']['notes'] for r in pair if r['supplement']['notes']]
                if len(defects)>4:notes.append('추가 불량: '+' / '.join(f"{source['part_no']} {defect['name']} {defect['qty']:g}" for source,defect in defects[4:]))
                put(sheet,'AG26','\n'.join(notes))
                for col,field in [('O','total_qty'),('Q','setup_qty'),('W','defect_qty')]:
                    vals=[r[field] for r in pair];put(sheet,col+'9',sum(vals) if all(v is not None for v in vals) else None)
                sheet.print_area='A1:AK34';sheet.sheet_properties.pageSetUpPr.fitToPage=True
                sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
                sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=1
                sheet.sheet_view.showGridLines=False
    workbook.remove(template)
    return workbook


def _color(color,default='#000000'):
    if color is not None and color.type=='rgb' and isinstance(color.rgb,str):
        return '#'+color.rgb[-6:]
    return default


def _border(side):
    if side is None or not side.style:return 'none'
    widths={'hair':'.35pt','thin':'.5pt','medium':'1pt','thick':'1.5pt','double':'1.5pt'}
    style='double' if side.style=='double' else ('dashed' if 'dash' in side.style.lower() else ('dotted' if side.style=='dotted' else 'solid'))
    return widths.get(side.style,'.5pt')+' '+style+' '+_color(side.color)


def workbook_html(workbook):
    parts=['<!doctype html><html lang="ko"><meta charset="utf-8"><title>작업일보</title><style>@page{size:A4 landscape;margin:10mm}*{box-sizing:border-box}body{margin:0;background:#e8edf3;font-family:"Malgun Gothic","맑은 고딕","Noto Sans CJK KR",sans-serif}.paper{width:297mm;min-height:210mm;padding:10mm;margin:18px auto;background:white;box-shadow:0 2px 10px #0002;break-after:page}.paper:last-child{break-after:auto}table{width:277mm;border-collapse:collapse;table-layout:fixed}td{padding:0 1px;overflow-wrap:anywhere;white-space:pre-wrap;line-height:1.1}.toolbar{position:sticky;top:0;padding:12px 20px;background:#fff;border-bottom:1px solid #cbd5e1;font-size:13px;z-index:1}.toolbar button{padding:8px 20px;background:#2563eb;color:white;border:0;border-radius:4px;margin-right:14px;cursor:pointer}@media print{body{background:white}.toolbar{display:none}.paper{width:277mm;min-height:0;padding:0;margin:0;box-shadow:none;print-color-adjust:exact;-webkit-print-color-adjust:exact}}</style><div class="toolbar"><button onclick="window.print()">인쇄</button> A4 가로 · 원본 양식 기준 · 미입력 값은 공란</div>']
    for sheet in workbook:
        merged={};skip=set()
        for area in sheet.merged_cells.ranges:
            merged[(area.min_row,area.min_col)]=(area.max_row-area.min_row+1,area.max_col-area.min_col+1)
            skip.update((r,c) for r in range(area.min_row,area.max_row+1) for c in range(area.min_col,area.max_col+1) if (r,c)!=(area.min_row,area.min_col))
        parts.append('<section class="paper"><table><colgroup>')
        widths=[]
        for c in range(1,38):
            dimension=next((d for d in sheet.column_dimensions.values() if (d.min or 0)<=c<=(d.max or 0)),None)
            widths.append(dimension.width if dimension else 8)
        for width in widths:parts.append(f'<col style="width:{width/sum(widths)*100:.4f}%">')
        parts.append('</colgroup>')
        for r in range(1,35):
            parts.append(f'<tr style="height:{sheet.row_dimensions[r].height or 14}pt">')
            for c in range(1,38):
                if (r,c) in skip:continue
                cell=sheet.cell(r,c);value=cell.value if cell.value is not None else ''
                if isinstance(value,float):value=f'{value:g}'
                rs,cs=merged.get((r,c),(1,1))
                title=html.escape(cell.comment.text,quote=True) if cell.comment else ''
                background=_color(cell.fill.fgColor,'white') if cell.fill.patternType=='solid' else 'white'
                font_size=cell.font.sz or 8
                # Blank spacer rows keep the template height without text line boxes.
                if value=='':font_size=0
                vertical={'center':'middle','top':'top','bottom':'bottom'}.get(cell.alignment.vertical,'middle')
                align=cell.alignment.horizontal or ('right' if isinstance(value,(float,int)) else 'left')
                styles=[f'font-size:{font_size}pt',f'font-weight:{"700" if cell.font.b else "400"}',
                        f'font-style:{"italic" if cell.font.i else "normal"}',f'color:{_color(cell.font.color)}',
                        f'background:{background}',f'text-align:{align}',f'vertical-align:{vertical}',
                        'white-space:'+('pre-wrap' if r!=34 and (cell.alignment.wrap_text or '\n' in str(value)) else 'pre')]
                for edge in ['left','right','top','bottom']:styles.append('border-'+edge+':'+_border(getattr(cell.border,edge)))
                style=html.escape(';'.join(styles),quote=True)
                parts.append(f'<td style="{style}" title="{title}" rowspan="{rs}" colspan="{cs}"><span class="cell-value">{html.escape(str(value))}</span></td>')
            parts.append('</tr>')
        parts.append('</table></section>')
    parts.append('''<script>
    document.fonts.ready.then(()=>requestAnimationFrame(()=>{
      document.querySelectorAll('.cell-value').forEach(span=>{
        const cell=span.parentElement;
        if(!span.textContent || getComputedStyle(cell).whiteSpace!=='pre')return;
        let size=parseFloat(getComputedStyle(cell).fontSize);
        const available=Math.max(1,cell.clientWidth-2);
        while(span.getBoundingClientRect().width>available && size>6){
          size-=.25;cell.style.fontSize=size+'px';
        }
      });
    }));
    </script></html>''');return ''.join(parts)
