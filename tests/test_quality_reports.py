import copy
from io import BytesIO
from types import SimpleNamespace
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from openpyxl import load_workbook
import models
from core.database import Base, get_db
from core.security import get_current_user
from models.models import ItemMasterModel, UserModel
from models.inspection_standard import InspectionStandard, InspectionStandardItem
from models.quality_report import InspectionReportMaster
from routers import quality_reports as routes
from services.quality_report_service import evaluate, build_report_excel, TEMPLATE_PATH


@pytest.fixture
def context(monkeypatch):
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    user = UserModel(username='quality', name='검사자', password_hash='x', role='ADMIN', created_at='2026-10-16')
    item = ItemMasterModel(part_no='TEST-A', part_name='시험품', account_type='PROD', material_type='FINISHED', created_at='2026-10-16')
    db.add_all([user, item]); db.flush()
    standard = InspectionStandard(item_id=item.id, document_type='FINAL', revision='02', status='CURRENT',
        part_no_snapshot=item.part_no, part_name_snapshot=item.part_name)
    standard.items = [InspectionStandardItem(sort_order=1, inspection_item_name='외경', spec_text='10 ±1', lower_limit=9, upper_limit=11, inspection_tool='V/C'),
                      InspectionStandardItem(sort_order=2, inspection_item_name='외관', spec_text='찍힘 없음')]
    db.add(standard); db.commit()
    app = FastAPI(); app.include_router(routes.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    payload = {'part_no': item.part_no, 'lot_no': 'LOT-001', 'inspection_date':'2026-10-16', 'standard_id':standard.id,
               'items':[{'standard_item_id':standard.items[0].id, 'values':[9,10,11,10,10]},
                        {'standard_item_id':standard.items[1].id, 'values':[None]*5, 'manual_judgment':'OK'}]}
    yield db, TestClient(app), payload, standard, user
    db.close(); engine.dispose()


def test_individual_ng_even_when_average_is_in_spec():
    average, judgment, samples = evaluate([8,10,10,10,12],9,11)
    assert average == 10 and judgment == 'NG'
    assert samples == ['NG','OK','OK','OK','NG']
    assert evaluate([9,10,None,10,11],9,11)[1] == 'HOLD'
    assert evaluate([None]*5,None,None)[1] == 'HOLD'
    assert evaluate([10]*5,10,None)[1] == 'OK'
    assert evaluate([1e308]*5, None, 1e308)[0] == 1e308


def test_save_snapshots_standard_and_preserves_excel_format(context):
    db, client, payload, standard, user = context
    response = client.post('/api/quality/outbound-reports',json=payload)
    assert response.status_code == 200, response.text
    result = response.json()['item']
    assert result['judgment'] == 'PASS' and result['inspector_name'] == '검사자'
    standard.items[0].spec_text = '변경된 규격'; standard.items[0].upper_limit = 5; db.commit()
    report = db.get(InspectionReportMaster, result['id'])
    assert report.items[0].spec_display == '10 ±1' and report.items[0].spec_max == 11
    template = load_workbook(TEMPLATE_PATH)
    exported = load_workbook(build_report_excel(report))
    sheet = exported.active
    assert sheet['C4'].value == 'TEST-A' and sheet['G4'].value == 'LOT-001'
    assert sheet['G10'].value == 9 and sheet['K10'].value == 10
    assert sheet['L10'].value == 10 and sheet['M10'].value == 'OK' and sheet['L38'].value == '합격'
    assert str(sheet.merged_cells) == str(template.active.merged_cells)
    assert sheet['G10']._style == template.active['G10']._style
    assert sheet.column_dimensions['E'].width == template.active.column_dimensions['E'].width
    assert sheet.page_setup.fitToHeight == 1
    assert client.get(f"/api/quality/outbound-reports/{report.id}/excel").content[:2] == b'PK'
    printed = client.get(f"/api/quality/outbound-reports/{report.id}/print")
    assert printed.status_code == 200 and 'LOT-001' in printed.text and '10 ±1' in printed.text


def test_missing_samples_hold_and_any_ng_rejects(context):
    _,client,payload,_,_ = context
    payload['items'][0]['values'][2] = None
    assert client.post('/api/quality/outbound-reports',json=payload).json()['item']['judgment'] == 'HOLD'
    payload['items'][0]['values'] = [8,10,10,10,12]
    assert client.post('/api/quality/outbound-reports',json=payload).json()['item']['judgment'] == 'REJECT'


def test_cannot_omit_duplicate_or_inject_standard_items(context):
    _,client,payload,_,_ = context
    invalid = copy.deepcopy(payload); invalid['items'].pop()
    assert client.post('/api/quality/outbound-reports',json=invalid).status_code == 422
    invalid = copy.deepcopy(payload); invalid['items'][1] = invalid['items'][0]
    assert client.post('/api/quality/outbound-reports',json=invalid).status_code == 422
    invalid = copy.deepcopy(payload); invalid['lot_no'] = '  '
    assert client.post('/api/quality/outbound-reports',json=invalid).status_code == 422
    invalid = copy.deepcopy(payload); invalid['items'][0]['values'][0] = 'NaN'
    assert client.post('/api/quality/outbound-reports',json=invalid).status_code == 422


def test_api_authorization_does_not_depend_on_referer(context, monkeypatch):
    _,client,payload,_,user = context
    monkeypatch.setattr(routes,'DEV_BYPASS_AUTH',False)
    user.role='USER'; user.permissions='{"menu_access":{"/quality/outbound-reports":"READ"}}'
    assert client.post('/api/quality/outbound-reports',json=payload).status_code == 403
    assert client.get('/api/quality/outbound-reports').status_code == 200
    user.permissions='{}'
    assert client.get('/api/quality/outbound-reports').status_code == 403


def test_text_in_excel_cannot_become_formula(context):
    db,client,payload,_,_ = context
    payload['note']='=HYPERLINK("https://example.invalid")'
    report_id = client.post('/api/quality/outbound-reports',json=payload).json()['item']['id']
    sheet = load_workbook(BytesIO(client.get(f'/api/quality/outbound-reports/{report_id}/excel').content)).active
    assert sheet['C39'].data_type == 's'


def test_all_27_items_export_through_row_36_and_28_are_rejected(context):
    db,client,payload,standard,_ = context
    for seq in range(3,28):
        source=InspectionStandardItem(sort_order=seq,inspection_item_name=f'치수 {seq}',spec_text='1~2',lower_limit=1,upper_limit=2)
        standard.items.append(source)
        db.flush()
        payload['items'].append({'standard_item_id':source.id,'values':[1.5]*5})
    db.commit()
    response=client.post('/api/quality/outbound-reports',json=payload)
    assert response.status_code==200,response.text
    report_id=response.json()['item']['id']
    sheet=load_workbook(BytesIO(client.get(f'/api/quality/outbound-reports/{report_id}/excel').content)).active
    assert sheet['B36'].value==27 and sheet['D36'].value=='치수 27'
    assert sheet['K36'].value==1.5 and sheet['M36'].value=='OK'
    payload['items'].append(payload['items'][-1])
    assert client.post('/api/quality/outbound-reports',json=payload).status_code==422
