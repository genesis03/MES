"""Generate the supplied default form; runtime exports only bind its cell values.

Run manually when intentionally updating the default layout. A company-approved
workbook may replace the asset if it retains the documented writable cells.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.worksheet.page import PageMargins


def build(path):
    book = Workbook()
    sheet = book.active
    sheet.title = '출하 검사 성적서'
    sheet.merge_cells('B1:M1')
    sheet['B1'] = '출하 검사 성적서'
    sheet['B1'].font = Font(name='맑은 고딕', size=20, bold=True)
    sheet['B1'].alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 36
    for col, width in {'A':2,'B':6,'C':12,'D':22,'E':24,'F':14,'G':11,'H':11,'I':11,'J':11,'K':11,'L':12,'M':10}.items():
        sheet.column_dimensions[col].width = width
    thin = Side(style='thin', color='333333')
    for row in sheet.iter_rows(min_row=3, max_row=39, min_col=2, max_col=13):
        for cell in row:
            cell.font = Font(name='맑은 고딕', size=9)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in range(3, 40):
        sheet.row_dimensions[r].height = 22 if r < 10 else 24
    for coord, value in {'B3':'관리번호','B4':'품번','B5':'품명','F3':'검사일','F4':'LOT','F5':'시료 수',
                         'K3':'작성','L3':'검토','M3':'승인','K38':'종합 판정','B39':'특기사항'}.items():
        sheet[coord] = value
    for address in ['C3:E3','C4:E4','C5:E5','G3:J3','G4:J4','G5:J5','L38:M38','C39:M39']:
        sheet.merge_cells(address)
    headers = ['순번','구분','검사항목','규격','계측기','x1','x2','x3','x4','x5','평균','판정']
    for column, value in enumerate(headers, 2):
        sheet.cell(9, column, value)
        sheet.cell(9, column).fill = PatternFill('solid', fgColor='E7EDF5')
        sheet.cell(9, column).font = Font(name='맑은 고딕', size=9, bold=True)
    sheet.merge_cells('B41:M41')
    sheet['B41'] = '(주)유니코어텍'
    sheet['B41'].alignment = Alignment(horizontal='center')
    sheet.print_area = 'B1:M41'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4, header=.1, footer=.1)
    sheet.print_options.horizontalCentered = True
    book.save(path)


if __name__ == '__main__':
    build(Path(__file__).resolve().parents[1] / 'assets' / 'InspectionReport_Template.xlsx')
