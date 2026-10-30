import json
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import models
from core.database import Base,get_db
from core.security import get_current_user
from models.models import ItemMasterModel,PurchaseInboundMaster,PurchaseInboundItem
from models.production import ProductionWorkOrder,ProductionPerformance
from models.production_lot import ProductionLotModel
from models.lot_consumption import LotConsumptionModel
from models.packing import PackingMaster,PackingBox,PackingLotAllocation
from models.sales import ShipmentMaster,ShipmentItem,ShipmentBox
from models.production_sync import ExternalProductionRecord
from models.packing_sync import ExternalPackingRecord,PackingSourcePresence
from models.subcontract_fault_sync import ExternalSubcontractFault
from models.shipment_sync import ExternalShipment,ExternalShipmentLot
from models.inventory_ledger import InventoryLedgerBalance,InventoryLedgerEntry,InventoryLedgerAudit
from routers.inventory_ledger import router
from services.inventory_ledger_service import report,detail


@pytest.fixture
def env(monkeypatch):
    from core import security
    monkeypatch.setattr(security,'DEV_BYPASS_AUTH',False)
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    Base.metadata.create_all(engine);db=sessionmaker(bind=engine)()
    items=[]
    for i,(part,kind,unit) in enumerate([('RAW','RAW','kg'),('SEMI-A','SEMI','EA'),('FIN','FINISHED','EA')],1):
        item=ItemMasterModel(id=i,part_no=part,part_name=part+' 내부품명',account_type='PROD',material_type=kind,unit=unit,created_at='2026-10-01',is_active='Y');items.append(item);db.add(item)
    db.commit()
    state=SimpleNamespace(user=SimpleNamespace(username='admin',role='ADMIN',permissions='{}'))
    app=FastAPI();app.include_router(router);app.dependency_overrides[get_db]=lambda:db;app.dependency_overrides[get_current_user]=lambda:state.user
    yield SimpleNamespace(db=db,client=TestClient(app),state=state,items=items)
    db.close();engine.dispose()


def baseline(env,item=2,**changes):
    payload={'month':'2026-10','source':'mes','item_id':item,'opening_qty':'0','actual_qty':None,'count_date':None,'note':'기초 등록','version':0}|changes
    response=env.client.put('/api/inventory/ledger/balance',json=payload)
    assert response.status_code==200,response.text
    return response


def event(env,qty='10',**changes):
    payload={'source':'mes','item_id':2,'entry_date':'2026-10-02','direction':'IN','category':'기타','qty':qty,'reason':'누락 거래 보완'}|changes
    r=env.client.post('/api/inventory/ledger/entries',json=payload)
    assert r.status_code==200,r.text
    return r.json()['id']


def row(env,item=2,source='mes'):
    return next(r for r in report(env.db,'2026-10',source)[0]['rows'] if r['item_id']==item)


def test_missing_opening_is_unknown_not_zero_and_zero_is_valid(env):
    event(env,'12')
    assert row(env)['closing_qty'] is None
    baseline(env)
    assert Decimal(row(env)['closing_qty'])==12
    assert Decimal(row(env)['opening_qty'])==0


def test_count_date_uses_that_day_not_month_end_and_no_stock_mutation(env):
    lot=ProductionLotModel(lot_no='L',item_id=2,part_no='SEMI-A',lot_qty=500);env.db.add(lot);env.db.commit()
    event(env,'10',entry_date='2026-10-02');event(env,'20',entry_date='2026-10-20')
    baseline(env,opening_qty='100',actual_qty='110',count_date='2026-10-10')
    value=row(env)
    assert Decimal(value['closing_qty'])==130
    assert Decimal(value['expected_at_count'])==110
    assert Decimal(value['difference'])==0
    assert env.db.get(ProductionLotModel,lot.id).lot_qty==500
    assert len(detail(env.db,'2026-10','mes',2)['daily'])==31


def test_optimistic_conflict_audit_and_blank_actual(env):
    baseline(env,actual_qty='0',count_date='2026-10-31')
    failed=env.client.put('/api/inventory/ledger/balance',json={'month':'2026-10','source':'mes','item_id':2,'opening_qty':'100','note':'덮어쓰기','version':0})
    assert failed.status_code==409
    env.db.rollback()
    assert Decimal(row(env)['opening_qty'])==0
    baseline(env,version=1,actual_qty=None,count_date=None,note='실사 입력 취소')
    assert row(env)['actual_qty'] is None
    assert env.db.query(InventoryLedgerAudit).filter_by(action='BALANCE').count()==2


@pytest.mark.parametrize('changes',[{'month':'2026-13'},{'actual_qty':'1'},{'count_date':'2026-09-30'},{'opening_qty':'NaN'},{'opening_qty':'-1'},{'opening_qty':'0.0000001'},{'note':' '}])
def test_invalid_balance_rejected(env,changes):
    response=env.client.put('/api/inventory/ledger/balance',json={'month':'2026-10','source':'mes','item_id':2,'opening_qty':'0','note':'실사','version':0}|changes)
    assert response.status_code==422,response.text
    assert env.db.query(InventoryLedgerBalance).count()==0


def test_invalid_query_returns_422(env):
    assert env.client.get('/api/inventory/ledger?month=2026-13').status_code==422
    assert env.client.get('/api/inventory/ledger?month=x').status_code==422


def test_void_is_audited_and_idempotent_and_preserves_entry(env):
    baseline(env);identifier=event(env,'0.125')
    assert Decimal(row(env)['closing_qty'])==Decimal('0.125')
    for _ in range(2):
        assert env.client.post(f'/api/inventory/ledger/entries/{identifier}/void',json={'reason':'중복 입력 정정'}).status_code==200
    assert Decimal(row(env)['closing_qty'])==0
    assert env.db.query(InventoryLedgerEntry).count()==1
    assert env.db.query(InventoryLedgerAudit).filter_by(action='VOID').count()==1


def test_read_only_and_missing_permissions_are_enforced_on_api(env):
    env.state.user=SimpleNamespace(username='reader',role='USER',permissions=json.dumps({'menu_access':{'/inventory/ledger':'READ'}}))
    assert env.client.get('/api/inventory/ledger?month=2026-10').status_code==200
    assert env.client.post('/api/inventory/ledger/entries',json={'source':'mes','item_id':2,'entry_date':'2026-10-01','direction':'OUT','category':'기타','qty':'1','reason':'사유'}).status_code==403
    env.state.user.permissions='{}'
    assert env.client.get('/api/inventory/ledger?month=2026-10').status_code==403
    assert env.client.get('/api/inventory/ledger/export?month=2026-10').status_code==403


def test_production_consumption_packing_shipping_count_once_by_business_date(env):
    db=env.db
    order=ProductionWorkOrder(work_order_no='WO1',order_date='2026-09-30',item_id=2,part_no='SEMI-A',order_qty=65)
    db.add(order);db.flush()
    perf=ProductionPerformance(work_order_id=order.id,performance_date='2026-10-02',process_code='LT',good_qty=60,defect_qty=3,setup_qty=2,created_at=datetime(2026,11,1))
    db.add(perf);db.flush()
    db.add(LotConsumptionModel(lot_no='RAWLOT',item_id=1,part_no='RAW',work_order_id=order.id,performance_id=perf.id,process_code='LT',consumed_qty=Decimal('6.5')))
    db.add(ProductionLotModel(lot_no='SEMI-LOT',item_id=2,part_no='SEMI-A',lot_qty=60,note=f'PERF:{perf.id}|실적'))
    pack=PackingMaster(packing_no='P1',packing_date='2026-10-03',item_id=3,part_no='FIN',box_count=1,box_qty=40,total_qty=40)
    pack.allocations=[PackingLotAllocation(source_lot_no='SEMI-LOT',allocated_qty=40)]
    pack.boxes=[PackingBox(box_no=1,package_lot_no='PACKLOT',box_qty=40)]
    db.add(pack);db.flush()
    shipment=ShipmentMaster(shipment_no='S1',shipment_date='2026-10-04',sales_order_id=1,customer_id=1,customer_name='고객')
    line=ShipmentItem(sales_order_item_id=1,item_id=3,part_no='FIN',shipped_qty=10)
    line.boxes=[ShipmentBox(packing_box_id=pack.boxes[0].id,package_lot_no='PACKLOT',shipped_qty=10)]
    shipment.items=[line];db.add(shipment);db.commit()
    baseline(env,item=1,opening_qty='10');baseline(env,item=2);baseline(env,item=3)
    assert Decimal(row(env,1)['closing_qty'])==Decimal('3.5')
    assert Decimal(row(env,2)['in_qty'])==65
    assert Decimal(row(env,2)['out_qty'])==45
    assert Decimal(row(env,2)['closing_qty'])==20
    assert Decimal(row(env,3)['closing_qty'])==30
    assert detail(db,'2026-10','mes',2)['events'][0]['date']=='2026-10-02'
    shipment.status='CANCELLED';pack.status='CANCELLED';db.commit()
    assert Decimal(row(env,2)['closing_qty'])==60
    assert Decimal(row(env,3)['closing_qty'])==0


def external_meta():
    return dict(content_hash='x',first_seen_at=datetime(2026,10,1),last_seen_at=datetime(2026,10,1),changed_at=datetime(2026,10,1))


def test_external_separate_from_native_no_packing_shipment_double_count(env):
    db=env.db
    packed=ExternalPackingRecord(source_key='P',part_no='FIN',lot_no='LOT',packing_date='2026-10-02',packing_qty='100',shipment_qty='40',shipment_date='2026-10-03',customer_name='고객',raw_json='{}',**external_meta())
    shipment=ExternalShipment(source_key='S',shipment_no='S1',shipment_date='2026-10-03',customer_name='고객',status='OUT_FIN',present=True,raw_json='{}',**external_meta())
    db.add_all([packed,shipment]);db.flush()
    db.add(ExternalShipmentLot(shipment_id=shipment.id,detail_key='D',source_lot_key='L',part_no='FIN',lot_no='LOT',allocated_qty='40'));db.commit()
    baseline(env,item=3,source='external')
    assert Decimal(row(env,3,'external')['closing_qty'])==60
    assert row(env,3,'mes')['event_count']==0
    db.add(PackingSourcePresence(record_id=packed.id,present=False,checked_at=datetime(2026,10,30)));shipment.present=False;db.commit()
    assert Decimal(row(env,3,'external')['closing_qty'])==0


def test_source_balance_isolation_and_current_master_name(env):
    baseline(env,opening_qty='10');baseline(env,source='external',opening_qty='20')
    assert Decimal(row(env)['closing_qty'])==10
    assert Decimal(row(env,source='external')['closing_qty'])==20
    env.items[1].part_name='新 내부품명';env.db.commit()
    assert row(env)['part_name']=='新 내부품명'


def test_export_filters_and_formula_injection_and_numeric_quantities(env):
    from openpyxl import load_workbook
    env.items[1].part_name='=HYPERLINK("bad")';env.db.commit()
    baseline(env,opening_qty='12.5',note='=SUM(1,2)');event(env,document_no='=1+1')
    response=env.client.get('/api/inventory/ledger/export?month=2026-10&part_no=SEMI-A')
    assert response.status_code==200
    book=load_workbook(BytesIO(response.content))
    assert book.worksheets[0].max_row==4
    assert book['실사 상세']['B2'].data_type=='s'
    assert book.worksheets[0]['D4'].value==12.5
    assert book['수불 상세']['G2'].data_type=='s'
    assert book['실사 상세']['I2'].data_type=='s'
    assert 'F2:H2' in book.worksheets[0].merged_cells
    assert book.worksheets[0].max_column==19
    assert book.worksheets[0]['P1'].value=='기말(산출)'
    assert book.worksheets[0]['Q1'].value=='재고실사'
    assert all(cell.value!='기말(전산)' for row in book.worksheets[0] for cell in row)
    assert book['일별 수불'].max_row==32
    assert book['일별 수불']['G32'].value==22.5
    assert '유형별 집계' in book.sheetnames


def test_external_production_and_subcontract_losses_are_separate_events(env):
    db=env.db
    production=ExternalProductionRecord(source_key='PROD',lot_no='A',job_no='JOB',part_no='SEMI-A',process_name='복합선반',work_date='2026-10-01',job_qty='60',lot_qty='65',fault_qty='3',raw_json=json.dumps({'F10':'2'}),**external_meta())
    # Exact manual connection prevents making an assumption based only on a suffix.
    from models.production_sync import ProductionSyncItemMap
    db.add(ProductionSyncItemMap(source_part_no='SOURCE',source_process_name='은도금',item_id=2))
    fault=ExternalSubcontractFault(source_key='FAULT',confirm_date='2026-10-02',inbound_no='I1',part_no='SOURCE',lot_no='B',process_name='은도금',defect_qty='25',present=True,raw_json=json.dumps({'ORI_LOT_QTY':'89','REST_QTY':'25','KIPUP_QTY':'4','LOSS_QTY':'0','SAMPLE_QTY':'0'}),**external_meta())
    db.add_all([production,fault]);db.commit()
    baseline(env,source='external')
    value=row(env,source='external')
    assert Decimal(value['in_qty'])==154
    assert Decimal(value['out_qty'])==34
    assert Decimal(value['closing_qty'])==120
    assert Decimal(value['production_qty'])==65
    assert Decimal(value['production_defect_qty'])==3
    assert Decimal(value['setup_qty'])==2
    assert Decimal(value['subcontract_in_qty'])==89
    assert Decimal(value['subcontract_defect_qty'])==25
    categories={e['category'] for e in detail(db,'2026-10','external',2)['events']}
    assert {'생산','생산불량','SET-UP','외주입고','외주불량','킵업'}<=categories
    assert report(db,'2026-10','external')[0]['warnings']
    fault.present=False;db.commit()
    assert Decimal(row(env,source='external')['closing_qty'])==60


def test_native_subcontract_outbound_and_inbound_change_their_own_part(env):
    from test_subcontract_fifo import dispatch,receive
    # Existing helper dispatches item 1 -> item 2, same warehouse movement as real routes.
    outbound=dispatch(env.db,'O1','2026-10-02',{'RAWLOT':100})
    inbound=receive(env.db,outbound,{outbound.items[0].lots[0].id:80})
    baseline(env,item=1,opening_qty='200');baseline(env,item=2)
    assert Decimal(row(env,1)['closing_qty'])==100
    assert Decimal(row(env,2)['closing_qty'])==80
    inbound.status='CANCELLED';outbound.status='CANCELLED';env.db.commit()
    assert Decimal(row(env,1)['closing_qty'])==200
    assert Decimal(row(env,2)['closing_qty'])==0


def test_inventory_adjustment_is_business_day_and_warehouse_transfer_is_not_consumption(env):
    from models.inventory_adjustment import InventoryAdjustmentModel
    from models.inventory_movement import InventoryMovementModel
    env.db.add(InventoryAdjustmentModel(item_id=2,part_no='SEMI-A',lot_no='L',before_qty=10,adjustment_qty=-2,after_qty=8,reason='차이',created_at=datetime(2026,10,31,23,59)))
    env.db.add(InventoryMovementModel(item_id=2,part_no='SEMI-A',lot_no='L',from_location='A',to_location='B',moved_qty=8,reason='이동',created_at=datetime(2026,10,31)))
    env.db.commit();baseline(env,opening_qty='10')
    assert Decimal(row(env)['closing_qty'])==8
    assert row(env)['event_count']==1


def test_unlinked_external_part_is_visible_but_name_is_never_taken_from_source(env):
    env.db.add(ExternalPackingRecord(source_key='UNLINKED',part_no='UNKNOWN',lot_no='LOT',packing_date='2026-10-01',packing_qty='30',shipment_qty='0',shipment_date='',customer_name='고객',raw_json=json.dumps({'PRODUCT_NM':'외부 이름'}),**external_meta()));env.db.commit()
    data=env.client.get('/api/inventory/ledger?month=2026-10&source=external&part_no=UNKNOWN').json()
    value=data['rows'][0]
    assert not value['linked'] and value['part_name']=='' and value['closing_qty'] is None
    assert env.client.get('/api/inventory/ledger/detail?month=2026-10&source=external&raw_part=UNKNOWN&process=포장').json()['events'][0]['qty']=='30'


def test_entry_direction_and_nonfinite_quantity_are_rejected(env):
    data={'source':'mes','item_id':2,'entry_date':'2026-10-01','direction':'OUT','category':'구매입고','qty':'1','reason':'누락'}
    assert env.client.post('/api/inventory/ledger/entries',json=data).status_code==422
    data.update(direction='IN',qty='Infinity')
    assert env.client.post('/api/inventory/ledger/entries',json=data).status_code==422


def test_count_carry_quantity_matches_workbook_difference(env):
    baseline(env,opening_qty='100',actual_qty='80',count_carry_qty='15',count_date='2026-10-31')
    assert Decimal(row(env)['difference'])==-5
    assert Decimal(row(env)['closing_qty'])==100


def batch(env,rows,**changes):
    return env.client.put('/api/inventory/ledger/balances',json={'month':'2026-10','source':'mes','note':'여러 품번 입력','rows':rows}|changes)


def test_bulk_only_entered_fields_change_and_zero_is_saved(env):
    baseline(env,opening_qty='100',actual_qty='50',count_carry_qty='3',count_date='2026-10-15')
    response=batch(env,[{'item_id':2,'version':1,'opening_qty':'0'}, {'item_id':3,'version':0,'actual_qty':'0'}],count_date='2026-10-31')
    assert response.status_code==200,response.text
    assert response.json()['saved']==2
    value=row(env)
    assert Decimal(value['opening_qty'])==0
    assert Decimal(value['actual_qty'])==50 and Decimal(value['count_carry_qty'])==3
    assert value['count_date']=='2026-10-15'  # Opening-only edits do not replace the previous inspection date.
    value=row(env,3)
    assert value['opening_qty'] is None and Decimal(value['actual_qty'])==0
    assert value['count_date']=='2026-10-31'
    assert env.db.query(InventoryLedgerAudit).count()==3
    assert row(env,2,'external')['opening_qty'] is None


def test_bulk_conflict_rolls_back_every_part_and_audit(env):
    baseline(env,item=2,opening_qty='10');baseline(env,item=3,opening_qty='20')
    response=batch(env,[{'item_id':2,'version':1,'opening_qty':'100'}, {'item_id':3,'version':0,'opening_qty':'200'}])
    assert response.status_code==409
    assert Decimal(row(env,2)['opening_qty'])==10
    assert Decimal(row(env,3)['opening_qty'])==20
    assert env.db.query(InventoryLedgerAudit).count()==2


def test_bulk_bad_merged_carry_rolls_back_previous_rows(env):
    baseline(env,item=2,opening_qty='10')
    response=batch(env,[{'item_id':2,'version':1,'opening_qty':'100'}, {'item_id':3,'version':0,'count_carry_qty':'5'}])
    assert response.status_code==422
    assert Decimal(row(env,2)['opening_qty'])==10
    assert env.db.query(InventoryLedgerBalance).count()==1
    assert env.db.query(InventoryLedgerAudit).count()==1


@pytest.mark.parametrize('rows,extra',[
    ([{'item_id':2,'version':0}],{}),
    ([{'item_id':2,'version':0,'opening_qty':None}],{}),
    ([{'item_id':2,'version':0,'actual_qty':'1'}],{}),
    ([{'item_id':2,'version':0,'actual_qty':'1'}],{'count_date':'2026-09-30'}),
    ([{'item_id':2,'version':0,'opening_qty':'NaN'}],{}),
    ([{'item_id':2,'version':0,'opening_qty':'1'},{'item_id':2,'version':0,'opening_qty':'2'}],{}),
])
def test_invalid_bulk_requests_do_not_write(env,rows,extra):
    assert batch(env,rows,**extra).status_code==422
    assert env.db.query(InventoryLedgerBalance).count()==0
    assert env.db.query(InventoryLedgerAudit).count()==0


def test_bulk_permissions_are_checked_without_relying_on_page_header(env):
    env.state.user=SimpleNamespace(username='reader',role='USER',permissions=json.dumps({'menu_access':{'/inventory/ledger':'READ'}}))
    assert batch(env,[{'item_id':2,'version':0,'opening_qty':'1'}]).status_code==403


def test_bulk_unknown_part_does_not_save_valid_preceding_part(env):
    response=batch(env,[{'item_id':2,'version':0,'opening_qty':'1'},{'item_id':999,'version':0,'opening_qty':'2'}])
    assert response.status_code==404
    assert env.db.query(InventoryLedgerBalance).count()==0
    assert env.db.query(InventoryLedgerAudit).count()==0
