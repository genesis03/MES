"""Control plan imports and revision persistence use isolated temporary databases."""
import io
import os
import tempfile
from pathlib import Path
from types import SimpleNamespace
import pytest
from openpyxl import Workbook
from services.control_plan_excel import HEADERS, parse_control_plan, ControlPlanImportError

_bootstrap=tempfile.TemporaryDirectory()
os.environ.setdefault('DATABASE_URL','sqlite:///'+str(Path(_bootstrap.name)/'bootstrap.db'))
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine,select,func
from sqlalchemy.orm import sessionmaker
from core.database import Base,get_db
from core.security import get_current_user
from models.models import CommonCodeModel,ItemMasterModel,UserModel
from models.process_flow import ProcessFlowRevision,ProcessFlowStep,ProcessFlowStepKey
from models.control_plan import ControlPlanRevision
from routers.control_plan import router

def workbook():
    book=Workbook();sheet=book.active;sheet.title='관리계획서'
    for cell,label in HEADERS.items(): sheet[cell]=label
    for r,number,name in [(14,10,'입고검사\n원자재'),(15,20,'가공')]:
        for c,value in {1:number,3:'◇',5:name,7:1,8:'치수',13:'10±1',14:'V/C',15:'n=3',16:'LOT',17:'검사일지',21:'●',23:'격리'}.items():sheet.cell(r,c,value)
        sheet.merge_cells(start_row=r,start_column=17,end_row=r,end_column=18)
        sheet.merge_cells(start_row=r,start_column=23,end_row=r,end_column=25)
        sheet.merge_cells(start_row=r,start_column=26,end_row=r,end_column=27)
    sheet['A2']='상단 회사명은 불러오지 않음';sheet['AB14']='인쇄 영역 밖 메모'
    return book

def data(book):
    result=io.BytesIO();book.save(result);return result.getvalue()

def steps():return [SimpleNamespace(id=1,step_no='10',step_name='입고검사'),SimpleNamespace(id=2,step_no='20',step_name='가공')]

def test_import_reference_formulas_and_item_metadata():
    book=workbook();sheet=book.active;sheet['M15']='=M14';sheet['G15']='=G14+1';sheet['H15']='=IF(H14="","",H14)'
    rows=parse_control_plan(data(book),steps())
    assert len(rows)==2
    assert rows[0]['process_detail']=='입고검사\n원자재'
    assert rows[1]['specification']=='10±1' and rows[1]['item_no']=='2' and rows[1]['product']=='치수'
    assert rows[0]['quality'] is True
    assert 'company' not in rows[0]

@pytest.mark.parametrize('cell,value',[('A15',30),('E15','다른 공정'),('A11','다른 항목'),('M15','=SUM(1,2)'),('M15','=M15'),('U15','예'),('H15','x'*4001)])
def test_import_rejects_mismatch_without_partial_result(cell,value):
    book=workbook();book.active[cell]=value
    with pytest.raises(ControlPlanImportError):parse_control_plan(data(book),steps())

def test_import_rejects_missing_and_reordered_processes():
    with pytest.raises(ControlPlanImportError):parse_control_plan(data(workbook()),list(reversed(steps())))
    with pytest.raises(ControlPlanImportError):parse_control_plan(data(workbook()),steps()+[SimpleNamespace(id=3,step_no='30',step_name='출하')])

def test_merged_number_can_have_distinct_characteristics():
    book=workbook();sheet=book.active
    sheet.insert_rows(15)
    for c,v in {1:10,3:'◇',5:'입고검사\n원자재',7:1,8:'치수2',13:'20±1',14:'V/C',17:'검사일지',23:'격리'}.items():sheet.cell(15,c,v)
    sheet.merge_cells('A14:A15');sheet.merge_cells('E14:E15');sheet.merge_cells('G14:G15')
    rows=parse_control_plan(data(book),steps())
    assert rows[0]['item_no']==rows[1]['item_no']=='1'
    assert rows[0]['product']=='치수' and rows[1]['product']=='치수2'

@pytest.fixture
def api(tmp_path):
    engine=create_engine('sqlite:///'+str(tmp_path/'cp.db'),connect_args={'check_same_thread':False})
    Base.metadata.create_all(engine);factory=sessionmaker(bind=engine,expire_on_commit=False)
    with factory() as db:
        user=UserModel(id=1,username='test',password_hash='not-login',name='테스터',role='ADMIN',created_at='2026-10-06')
        db.add(user);db.add(CommonCodeModel(group_code='MATERIAL_TYPE',code='FINISHED',group_name='재질',code_name='완제품',created_at='2026-10-06'))
        db.add(ItemMasterModel(id=1,part_no='310186-1',part_name='회사 완제품',account_type='PROD',material_type='FINISHED',created_at='2026-10-06'))
        db.add(ProcessFlowRevision(id=1,item_id=1,revision_code='REV.0',sequence=1,version=1,status='CURRENT',part_no_snapshot='310186-1',part_name_snapshot='회사 완제품',created_by_id=1,created_by='테스터'));db.flush()
        for step in steps():
            db.add(ProcessFlowStepKey(id=step.id,item_id=1));db.flush()
            db.add(ProcessFlowStep(id=step.id,revision_id=1,step_key_id=step.id,sort_order=step.id,step_no=step.step_no,step_name=step.step_name))
        db.commit()
    app=FastAPI();app.include_router(router)
    def sessions():
        with factory() as db:yield db
    app.dependency_overrides[get_db]=sessions;app.dependency_overrides[get_current_user]=lambda:user
    with TestClient(app) as client:yield client,factory,user
    engine.dispose()

def upload(client,book=None,**extra):
    return client.post('/api/control-plans/import-excel',data={'item_id':1,'flow_revision_id':1,'flow_version':1,**extra},files={'file':('company.xlsx',data(book or workbook()))})

def payload():
    return dict(item_id=1,document_no='D-110-008-005',revision_code='REV.0',flow_revision_id=1,flow_version=1,
                header={'company':'회사명 유지','stage_production':True,'prepared_signature':'작성자'},rows=parse_control_plan(data(workbook()),steps()))

def test_validate_then_save_activate_revise_and_keep_history(api):
    client,factory,_=api
    assert upload(client).status_code==200
    with factory() as db:assert db.scalar(select(func.count()).select_from(ControlPlanRevision))==0
    body=payload();created=client.post('/api/control-plans/revisions',json=body);assert created.status_code==200,created.text
    saved=created.json();id=saved['id']
    body['version']=saved['version'];body['header']['company']='수정 회사명'
    saved=client.put(f'/api/control-plans/revisions/{id}',json=body).json()
    assert saved['header']['company']=='수정 회사명'
    assert saved['header']['part_no']=='310186-1'
    assert client.put(f'/api/control-plans/revisions/{id}',json=body).status_code==409
    applied=client.post(f'/api/control-plans/revisions/{id}/activate',json={'version':saved['version']});assert applied.status_code==200,applied.text
    body['version']=applied.json()['version']
    assert client.put(f'/api/control-plans/revisions/{id}',json=body).status_code==409
    revised=client.post(f'/api/control-plans/revisions/{id}/revise',json={'version':applied.json()['version'],'revision_code':'1'});assert revised.status_code==200,revised.text
    assert revised.json()['status']=='DRAFT' and revised.json()['rows']==applied.json()['rows']
    new=client.post(f"/api/control-plans/revisions/{revised.json()['id']}/activate",json={'version':1});assert new.status_code==200,new.text
    old=client.get(f'/api/control-plans/revisions/{id}').json()
    assert old['status']=='SUPERSEDED' and old['rows']==applied.json()['rows']

def test_failed_import_keeps_saved_content_and_version(api):
    client,_,_=api;saved=client.post('/api/control-plans/revisions',json=payload()).json()
    book=workbook();book.active['M14']='=UNSUPPORTED()'
    response=upload(client,book,revision_id=saved['id'],revision_version=saved['version']);assert response.status_code==422
    assert client.get('/api/control-plans/revisions/'+str(saved['id'])).json()==saved

def test_guards_permissions_flow_version_identity_and_order(api):
    client,factory,user=api
    body=payload();body['rows'].reverse()
    assert client.post('/api/control-plans/revisions',json=body).status_code==422
    assert upload(client,flow_version=2).status_code==409
    user.role='USER';user.permissions='{"menu_access":{"/standard-documents/control-plans":"READ"}}'
    assert client.get('/api/control-plans/options').status_code==200
    assert upload(client).status_code==403
    assert client.post('/api/control-plans/revisions',json=payload()).status_code==403
    user.permissions='{}'
    assert client.get('/api/control-plans/revisions?item_id=1').status_code==403

def test_canonical_revision_and_snapshot_survive_source_change(api):
    client,factory,_=api;body=payload();body['revision_code']='00'
    saved=client.post('/api/control-plans/revisions',json=body).json();assert saved['revision_code']=='REV.0'
    assert client.post('/api/control-plans/revisions',json=payload()).status_code==409
    with factory() as db:
        db.get(ProcessFlowStep,1).step_name='변경 공정';db.get(ProcessFlowRevision,1).version=2;db.commit()
    read=client.get('/api/control-plans/revisions/'+str(saved['id'])).json()
    assert read['flow']['steps'][0]['step_name']=='입고검사'
    assert client.post('/api/control-plans/revisions/'+str(saved['id'])+'/activate',json={'version':1}).status_code==409
