import io
from openpyxl import load_workbook
from test_production_sync import setup,DAY,row
from routers.daily_job_report import router,Supplement
from services import production_sync_service as sync
from services.daily_job_report import build_workbook,workbook_html
from models.daily_job_report import DailyJobReportSupplement
from models.production_sync import ExternalProductionRecord
from models.production_lot import ProductionLotModel

def cell(sheet,coordinate):
    return sheet[sheet._report_cell_map.get(coordinate,coordinate)]



def test_external_supplement_round_trip_and_readonly_export(setup):
    setup.app.include_router(router)
    token,_=sync.claim(DAY,DAY,'MANUAL');sync.save_day([row()],token)
    result=setup.client.get('/api/production/daily-job-report',params={'day':DAY.isoformat()})
    assert result.status_code==200,result.text
    source=result.json()['items'][0];key=source['key']
    supplement={'raw_input':5,'raw_remaining':2,'cut_bars':1,'raw_setup':7,'operator_name':'Operator','shift':'야간',
                'measurements':[{'label':'Length','spec':'10±0.1','values':[9.99,10,10.01]}]}
    url='/api/production/daily-job-report/EXTERNAL/'+str(source['source_record_id'])
    assert setup.client.put(url,params={'day':DAY.isoformat()},json=supplement).status_code==200
    persisted=setup.client.get('/api/production/daily-job-report',params={'day':DAY.isoformat()}).json()['items'][0]
    assert persisted['supplement']['measurements'][0]['values']==[9.99,10,10.01]
    payload={'day':DAY.isoformat(),'keys':[key]}
    response=setup.client.post('/api/production/daily-job-report/excel',json=payload)
    assert response.status_code==200,response.text
    sheet=load_workbook(io.BytesIO(response.content)).active
    sheet._report_cell_map=build_workbook([persisted]).active._report_cell_map
    assert [cell(sheet,x).value for x in ['D16','F16','G16']]==[9.99,10,10.01]
    assert cell(sheet,'Y12').value=='5 / 2' and cell(sheet,'AC12').value==1 and cell(sheet,'AF12').value==7
    assert cell(sheet,'C3').value=='야간' and cell(sheet,'O7').value==331 and cell(sheet,'Q7').value==115
    assert cell(sheet,'AA15').value is None
    html=setup.client.post('/api/production/daily-job-report/print',json=payload)
    assert html.status_code==200 and '측정 3' in html.text
    with setup.sessions() as db:
        assert db.query(DailyJobReportSupplement).count()==1
        assert db.query(ExternalProductionRecord).one().job_qty=='216'
        assert db.query(ProductionLotModel).count()==0
    assert setup.client.put(url,params={'day':DAY.isoformat()},json={'measurements':[{'values':[1]}]}).status_code==422
    assert setup.client.put(url,params={'day':'2026-01-01'},json=supplement).status_code==409


def report_row(lot,measurements=0):
    extra=Supplement(measurements=[{'label':str(i),'values':[i,i+.1,i+.2]} for i in range(measurements)]).model_dump()
    return dict(performance_date='2026-10-09',process_name='복합선반',equipment_name='1호기',shift_name='주간',
                operator_name='Worker',part_no='=1+1',part_name='Test',total_qty=100,setup_qty=2,defect_qty=1,
                work_time='20261009090000 ~ 20261009180000',material_lots='RAW: 1 KG',output_lot_no=lot,supplement=extra)


def test_multiple_work_rows_and_inspection_items_are_not_truncated():
    workbook=build_workbook([report_row('LOT1',10),report_row('LOT2'),report_row('LOT3')])
    assert len(workbook.worksheets)==3
    assert cell(workbook.worksheets[0],'AH7').value=='LOT1' and cell(workbook.worksheets[0],'AH8').value=='LOT2'
    assert cell(workbook.worksheets[1],'D16').value==8
    assert cell(workbook.worksheets[2],'AH7').value=='LOT3'
    assert cell(workbook.worksheets[0],'G7').data_type=='s'
    assert cell(workbook.worksheets[0],'M7').value=='09:00 ~ 18:00'
    assert 'LOT3' in workbook_html(workbook) and '검사 계속 2/2' in workbook_html(workbook)


def test_native_quantities_and_current_inspection_standard_match(setup):
    from test_production_sync import add_item
    from models.models import ProcessModel
    from models.production import ProductionWorkOrder,ProductionPerformance
    from models.inspection_standard import InspectionStandard,InspectionStandardItem
    setup.app.include_router(router)
    with setup.sessions() as db:
        item=add_item(db)
        db.add(ProcessModel(process_code='LT',process_name='복합선반',created_at='2026-01-01'))
        db.flush()
        order=ProductionWorkOrder(work_order_no='W1',order_date=DAY.isoformat(),item_id=item,part_no='LOCAL-B',order_qty=500)
        db.add(order);db.flush()
        db.add(ProductionPerformance(work_order_id=order.id,performance_date=DAY.isoformat(),process_code='LT',shift_type='DAY',equipment_name='복합선반 1호기',operator_name='Worker',good_qty=216,defect_qty=0,setup_qty=115))
        standard=InspectionStandard(document_type='PROCESS',item_id=item,inspection_process_code='LT',revision='Rev.00',status='CURRENT',part_no_snapshot='LOCAL-B',part_name_snapshot='Test')
        standard.items.append(InspectionStandardItem(inspection_item_name='외경',spec_text='Ø10±0.1'))
        db.add(standard);db.commit()
    data=setup.client.get('/api/production/daily-job-report',params={'day':DAY.isoformat()}).json()['items'][0]
    assert data['total_qty']==331 and data['setup_qty']==115 and data['operator_name']=='Worker'
    assert data['supplement']['measurements'][0]['label']=='외경'
    assert data['supplement']['measurements'][0]['values']==[None,None,None]


def test_print_preserves_original_fonts_fills_borders_and_blank_spacers():
    workbook=build_workbook([report_row('LOT1')])
    output=workbook_html(workbook)
    assert 'font-size:18.0pt' in output
    assert 'font-weight:700' in output and 'background:#E0E0E0' in output
    assert 'background:#FFFF00' not in output
    assert 'border-left:1pt solid' in output
    assert 'font-size:0pt' in output
    assert cell(workbook.active,'D7').alignment.horizontal=='left'
    assert cell(workbook.active,'F16').fill.patternType is None
    assert cell(workbook.active,'G14').font.sz==7


def test_measurement_widths_grow_without_moving_other_report_sections():
    from services.daily_job_report import _column_widths,TEMPLATE
    source=load_workbook(TEMPLATE).active
    target=build_workbook([report_row('LOT1',1)]).active
    widths=_column_widths(target);original=_column_widths(source)
    assert abs(sum(widths)-sum(original))<1e-5
    for coordinate in ['A1','C5','Z5','AH5','A26','B34']:
        before=source[coordinate];after=cell(target,coordinate)
        assert abs(sum(original[:before.column-1])-sum(widths[:after.column-1]))<1e-5
    def width(coordinate):
        c=cell(target,coordinate)
        area=next((a for a in target.merged_cells.ranges if c.coordinate in a),None)
        return sum(widths[area.min_col-1:area.max_col]) if area else widths[c.column-1]
    samples=[width(c) for c in ['D16','F16','G16']]
    assert max(samples)-min(samples)<1e-5
    assert samples[0]>sum(original[3:7])/3*2
    assert width('H16')<sum(original[7:9]) and width('J16')<original[9]
    assert all(c.fill.fgColor.rgb!='FFFFFF00' for row in target for c in row)


def test_processed_defects_link_to_exact_performance_and_exclude_cancelled(setup):
    from test_production_sync import add_item
    from models.models import CommonCodeModel, ProcessModel
    from models.production import ProductionWorkOrder, ProductionPerformance
    from models.production_defect import QualityProductionDefect, QualityProductionDefectDetail
    setup.app.include_router(router)
    with setup.sessions() as db:
        item=add_item(db)
        db.add(ProcessModel(process_code='LT',process_name='복합선반',created_at='2026-01-01'))
        db.flush()
        order=ProductionWorkOrder(work_order_no='W-DEFECT',order_date=DAY.isoformat(),item_id=item,part_no='LOCAL-B',order_qty=100)
        db.add(order);db.flush()
        performance=ProductionPerformance(work_order_id=order.id,performance_date=DAY.isoformat(),process_code='LT',shift_type='DAY',good_qty=100,defect_qty=0,setup_qty=0)
        db.add(performance);db.flush()
        db.add(CommonCodeModel(group_code='DEFECT_TYPE',group_name='보고서 테스트',created_at='2026-01-01',code='REPORT_TEST',code_name='치수 불량',is_active='N'))
        db.add(CommonCodeModel(group_code='PRODUCTION_DEFECT_REASON',group_name='보고서 테스트',created_at='2026-01-01',code='REPORT_TEST',code_name='검사 불량',is_active='N'))
        for number,status,marker in [(1,'ACTIVE',performance.id),(2,'CANCELLED',performance.id),(3,'ACTIVE',performance.id+100)]:
            lot=ProductionLotModel(lot_no=f'REPORTLOT{number}',item_id=item,part_no='LOCAL-B',lot_qty=100,note=f'PERF:{marker}|생산실적 자동생성')
            db.add(lot);db.flush()
            treatment=QualityProductionDefect(production_lot_id=lot.id,item_id=item,lot_no=lot.lot_no,defect_date=DAY.isoformat(),defect_qty=3,status=status,defect_reason_code='REPORT_TEST',remark='선별 후 폐기')
            db.add(treatment);db.flush()
            db.add(QualityProductionDefectDetail(defect_id=treatment.id,defect_type_code='REPORT_TEST',defect_qty=3))
        db.commit()
    response=setup.client.get('/api/production/daily-job-report',params={'day':DAY.isoformat()})
    assert response.status_code==200,response.text
    data=response.json()['items'][0]
    assert data['defects']==[{'name':'치수 불량','qty':3,'result':'검사 불량 / 선별 후 폐기'}]
    assert data['total_qty']==100 and data['defect_qty']==0
    sheet=build_workbook([data]).active
    assert cell(sheet,'AD15').value=='치수 불량'
    assert cell(sheet,'AJ15').value=='검사 불량 / 선별 후 폐기'
    assert cell(sheet,'AA15').value==3
    assert '선별 후 폐기' in workbook_html(build_workbook([data]))
